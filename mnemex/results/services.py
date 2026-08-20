from __future__ import annotations

import csv
import hashlib
import json
import unicodedata
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from enum import Enum
from io import BytesIO, StringIO
from pathlib import PurePath

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook  # type: ignore[import-untyped]

from mnemex.accounts.authorization import Action, may_perform
from mnemex.accounts.models import Account
from mnemex.foundation.models import IdempotencyRecord
from mnemex.foundation.services import claim_idempotency, record_audit_event
from mnemex.partners.models import PartnerOrganization
from mnemex.results.models import (
    IngestionRun,
    MappingTemplate,
    PublishedSourceResult,
    ReconciliationCase,
    SourceArtifact,
    StagedResult,
)
from mnemex.schema import Discipline, ScoreType

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_XLSX_ENTRIES = 100
MAX_ROWS = 2_000
MAX_COLUMNS = 100
MAX_CELL_CHARACTERS = 10_000
MAX_SOURCE_REVISION = 2_147_483_647
MAX_SCORE_SIGNIFICANT_DIGITS = 64
MAX_CANONICAL_SCORE_CHARACTERS = 128

_REQUIRED_FIELDS = (
    "source_result_id",
    "source_revision",
    "source_event_id",
    "event_name",
    "result_date",
    "competitor_name",
    "discipline",
    "score_type",
    "score",
)
_OPTIONAL_FIELDS = ("heat_id", "wood_species", "wood_diameter_mm", "wood_quality")
_CSV_TYPES = {"text/csv", "application/csv", "application/vnd.ms-excel"}
_XLSX_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",
}


@dataclass(frozen=True)
class IngestionOutcome:
    run: IngestionRun
    replayed: bool


@dataclass(frozen=True)
class ParsedSpreadsheet:
    rows: list[dict[str, object]]
    source_sheet_name: str
    source_row_numbers: list[int]


class ResultPreviewClassification(str, Enum):
    """Side-effect-free outcome names shared with migration previews."""

    PUBLISHABLE = "publishable"
    QUARANTINED = "quarantined"
    DUPLICATE = "duplicate"
    CONFLICT = "conflict"


@dataclass(frozen=True)
class ResultValidationPreview:
    """Canonical Results Desk validation and current-state classification."""

    normalized_payload: dict[str, object]
    payload_digest: str
    validation_errors: list[dict[str, str]]
    classification: ResultPreviewClassification


class NumericExpansionError(ValueError):
    """A finite Decimal cannot be safely represented in canonical fixed point."""


def _sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _canonical_bytes(value: object) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")


def _jsonb_safe_projection(value: object) -> object:
    """Return diagnostic data that PostgreSQL JSONB can persist safely.

    The source artifact digest remains derived from the original input. This
    projection only escapes characters that PostgreSQL JSONB cannot represent,
    so an invalid row can still be quarantined and inspected.
    """

    if isinstance(value, str):
        return "".join(
            f"\\u{ord(character):04x}"
            if character == "\x00" or 0xD800 <= ord(character) <= 0xDFFF
            else character
            for character in value
        )
    if isinstance(value, Mapping):
        return {
            str(_jsonb_safe_projection(key)): _jsonb_safe_projection(item)
            for key, item in value.items()
        }
    if isinstance(value, Sequence) and not isinstance(value, (bytes, bytearray)):
        return [_jsonb_safe_projection(item) for item in value]
    return value


def _authorize(actor: Account, organization: PartnerOrganization) -> None:
    if not may_perform(actor, Action.MANAGE_RESULTS, organization=organization):
        raise PermissionDenied("an MFA-bound results manager role is required")


def _validate_context(
    *, organization: PartnerOrganization, mapping_template: MappingTemplate
) -> MappingTemplate:
    current_organization = PartnerOrganization.objects.filter(pk=organization.pk).first()
    current_mapping = (
        MappingTemplate.objects.select_for_update().filter(pk=mapping_template.pk).first()
    )
    if (
        current_organization is None
        or current_organization.status != PartnerOrganization.Status.ACTIVE
    ):
        raise ValidationError("source organization is not active")
    if (
        current_mapping is None
        or not current_mapping.is_active
        or current_mapping.organization_id != current_organization.organization_id
    ):
        raise ValidationError("mapping template is not active for the source organization")
    if not isinstance(current_mapping.field_map, dict):
        raise ValidationError("mapping template field_map must be an object")
    missing = [field for field in _REQUIRED_FIELDS if not current_mapping.field_map.get(field)]
    if missing:
        raise ValidationError(f"mapping template is missing canonical fields: {', '.join(missing)}")
    return current_mapping


def _safe_cell(value: object) -> object:
    if value is None or isinstance(value, (str, int, float, bool)):
        result = value
    else:
        result = str(value)
    if isinstance(result, str) and len(result) > MAX_CELL_CHARACTERS:
        raise ValidationError("a spreadsheet cell exceeds the safe character limit")
    return result


def _validated_headers(raw_headers: Sequence[object], *, file_type: str) -> list[str]:
    headers = [str(value).strip() if value is not None else "" for value in raw_headers]
    if any(not header for header in headers):
        raise ValidationError(f"{file_type} header names cannot be blank")
    if len(set(headers)) != len(headers):
        raise ValidationError(f"{file_type} header names must be unique")
    return headers


def _canonical_score(score: Decimal) -> str:
    if score == 0:
        return "0"
    sign, decimal_digits, exponent = score.as_tuple()
    if not isinstance(exponent, int):
        raise NumericExpansionError
    significant_length = len(decimal_digits)
    while significant_length > 1 and decimal_digits[significant_length - 1] == 0:
        significant_length -= 1
    digits = decimal_digits[:significant_length]
    exponent += len(decimal_digits) - significant_length
    if len(digits) > MAX_SCORE_SIGNIFICANT_DIGITS:
        raise NumericExpansionError
    point = len(digits) + exponent
    if exponent >= 0:
        expanded_length = sign + len(digits) + exponent
    elif point > 0:
        expanded_length = sign + len(digits) + 1
    else:
        expanded_length = sign + 2 + (-point) + len(digits)
    if expanded_length > MAX_CANONICAL_SCORE_CHARACTERS:
        raise NumericExpansionError
    fixed = format(score, "f")
    if "." in fixed:
        fixed = fixed.rstrip("0").rstrip(".")
    return fixed


def _parse_csv(content: bytes) -> ParsedSpreadsheet:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationError("CSV files must use UTF-8 encoding") from exc
    reader = csv.DictReader(StringIO(text, newline=""))
    if reader.fieldnames is None:
        raise ValidationError("CSV file must contain a header row")
    if len(reader.fieldnames) > MAX_COLUMNS:
        raise ValidationError("spreadsheet exceeds the safe column limit")
    reader.fieldnames = _validated_headers(reader.fieldnames, file_type="CSV")
    rows: list[dict[str, object]] = []
    source_row_numbers: list[int] = []
    for data_row_number, row in enumerate(reader, start=1):
        if data_row_number > MAX_ROWS:
            raise ValidationError("spreadsheet exceeds the safe row limit")
        if None in row:
            raise ValidationError("CSV row contains more values than the header row")
        rows.append({str(key): _safe_cell(value) for key, value in row.items() if key is not None})
        source_row_numbers.append(data_row_number + 1)
    return ParsedSpreadsheet(
        rows=rows,
        source_sheet_name="",
        source_row_numbers=source_row_numbers,
    )


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise ValidationError("XLSX archive exceeds the safe entry limit")
            if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValidationError("XLSX archive exceeds the safe uncompressed-size limit")
    except zipfile.BadZipFile as exc:
        raise ValidationError("XLSX file is not a valid workbook archive") from exc


def _available_worksheets(sheet_names: Sequence[str]) -> str:
    shown = [repr(name) for name in sheet_names[:10]]
    suffix = f", plus {len(sheet_names) - 10} more" if len(sheet_names) > 10 else ""
    return f"{', '.join(shown)}{suffix}"


def _select_worksheet(sheet_names: Sequence[str], requested_name: str | None) -> str:
    requested = (requested_name or "").strip()
    available = _available_worksheets(sheet_names)
    if requested:
        if requested not in sheet_names:
            raise ValidationError(
                f"XLSX worksheet {requested!r} was not found. Available worksheets: {available}."
            )
        return requested
    if len(sheet_names) != 1:
        raise ValidationError(
            "XLSX workbook contains multiple worksheets; enter one worksheet name. "
            f"Available worksheets: {available}."
        )
    return sheet_names[0]


def _parse_xlsx(content: bytes, *, worksheet_name: str | None) -> ParsedSpreadsheet:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise ValidationError("XLSX file could not be parsed safely") from exc
    try:
        selected_sheet_name = _select_worksheet(workbook.sheetnames, worksheet_name)
        sheet = workbook[selected_sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        raw_headers = next(iterator, None)
        if raw_headers is None:
            raise ValidationError("XLSX file must contain a header row")
        if len(raw_headers) > MAX_COLUMNS:
            raise ValidationError("spreadsheet exceeds the safe column limit")
        headers = _validated_headers(raw_headers, file_type="XLSX")
        rows: list[dict[str, object]] = []
        source_row_numbers: list[int] = []
        for data_row_number, values in enumerate(iterator, start=1):
            if data_row_number > MAX_ROWS:
                raise ValidationError("spreadsheet exceeds the safe row limit")
            rows.append(
                {
                    header: _safe_cell(values[index] if index < len(values) else None)
                    for index, header in enumerate(headers)
                }
            )
            source_row_numbers.append(data_row_number + 1)
        return ParsedSpreadsheet(
            rows=rows,
            source_sheet_name=selected_sheet_name,
            source_row_numbers=source_row_numbers,
        )
    finally:
        workbook.close()


def _mapped_payload(
    row: Mapping[str, object], field_map: Mapping[str, str]
) -> tuple[dict[str, object], list[dict[str, str]]]:
    payload = {
        field: _safe_cell(row.get(source_field)) for field, source_field in field_map.items()
    }
    for field in _OPTIONAL_FIELDS:
        payload.setdefault(field, None)
    errors: list[dict[str, str]] = []

    def raw_text(field: str) -> str:
        value = payload.get(field)
        return str("" if value is None else value).strip()

    def plain_text(field: str, *, max_length: int, required_message: str | None = None) -> str:
        value = raw_text(field)
        payload[field] = value
        if not value and required_message is not None:
            errors.append({"field": field, "code": "required", "message": required_message})
        elif value and (
            len(value) > max_length
            or any(unicodedata.category(character).startswith("C") for character in value)
        ):
            errors.append(
                {
                    "field": field,
                    "code": "invalid_text",
                    "message": f"Enter plain text of {max_length} characters or fewer.",
                }
            )
        return value

    source_result_id = plain_text(
        "source_result_id", max_length=200, required_message="Enter a source result ID."
    )

    try:
        revision = int(raw_text("source_revision"))
        if revision < 1 or revision > MAX_SOURCE_REVISION:
            raise ValueError
        payload["source_revision"] = revision
    except ValueError:
        payload["source_revision"] = None
        errors.append(
            {
                "field": "source_revision",
                "code": "invalid_revision",
                "message": f"Enter a revision from 1 to {MAX_SOURCE_REVISION}.",
            }
        )

    plain_text("source_event_id", max_length=200, required_message="Enter a source event ID.")
    plain_text("event_name", max_length=240, required_message="Enter the event name.")

    result_date = raw_text("result_date")
    payload["result_date"] = result_date
    if not result_date:
        errors.append(
            {"field": "result_date", "code": "required", "message": "Enter a result date."}
        )
    else:
        try:
            parsed_date = date.fromisoformat(result_date)
            if parsed_date.isoformat() != result_date:
                raise ValueError
        except ValueError:
            errors.append(
                {
                    "field": "result_date",
                    "code": "invalid_date",
                    "message": "Enter the result date as YYYY-MM-DD.",
                }
            )

    plain_text(
        "competitor_name",
        max_length=240,
        required_message="Enter the source competitor name.",
    )
    plain_text("heat_id", max_length=200)
    plain_text("wood_species", max_length=160)

    discipline = str(payload.get("discipline") or "").strip().upper()
    payload["discipline"] = discipline
    if discipline not in {item.value for item in Discipline}:
        errors.append(
            {
                "field": "discipline",
                "code": "unknown_discipline",
                "message": "Choose a recognized MNEMEX discipline.",
            }
        )

    score_type = str(payload.get("score_type") or "").strip().lower()
    payload["score_type"] = score_type
    if score_type not in {item.value for item in ScoreType}:
        errors.append(
            {
                "field": "score_type",
                "code": "unknown_score_type",
                "message": "Choose a recognized score type.",
            }
        )

    try:
        score = Decimal(raw_text("score"))
        if not score.is_finite():
            raise InvalidOperation
        payload["score"] = _canonical_score(score)
        if score_type == ScoreType.TIME.value and score <= 0:
            errors.append(
                {
                    "field": "score",
                    "code": "must_be_positive",
                    "message": "Enter a time score greater than zero.",
                }
            )
        elif score_type in {item.value for item in ScoreType} and score < 0:
            errors.append(
                {
                    "field": "score",
                    "code": "must_be_non_negative",
                    "message": "Enter a score of zero or greater.",
                }
            )
    except NumericExpansionError:
        payload["score"] = raw_text("score")
        errors.append(
            {
                "field": "score",
                "code": "number_out_of_bounds",
                "message": (
                    f"Enter a score with at most {MAX_SCORE_SIGNIFICANT_DIGITS} significant "
                    f"digits and {MAX_CANONICAL_SCORE_CHARACTERS} fixed-point characters."
                ),
            }
        )
    except InvalidOperation:
        payload["score"] = raw_text("score")
        errors.append(
            {"field": "score", "code": "invalid_number", "message": "Enter a numeric score."}
        )

    def optional_integer(
        field: str, *, minimum: int, maximum: int | None, code: str, message: str
    ) -> None:
        raw = raw_text(field)
        if not raw:
            payload[field] = None
            return
        try:
            number = Decimal(raw)
            if not number.is_finite() or number != number.to_integral_value():
                raise InvalidOperation
            safe_maximum = maximum if maximum is not None else MAX_SOURCE_REVISION
            if number < minimum or number > safe_maximum:
                raise InvalidOperation
            value = int(number)
        except (InvalidOperation, ValueError, OverflowError):
            payload[field] = raw
            errors.append({"field": field, "code": code, "message": message})
        else:
            payload[field] = value

    optional_integer(
        "wood_diameter_mm",
        minimum=1,
        maximum=None,
        code="invalid_positive_integer",
        message="Enter a whole-number wood diameter greater than zero, or leave it blank.",
    )
    optional_integer(
        "wood_quality",
        minimum=1,
        maximum=10,
        code="out_of_range",
        message="Enter a wood quality from 1 to 10, or leave it blank.",
    )

    return payload, errors


def preview_mapped_result(
    *,
    organization: PartnerOrganization,
    row: Mapping[str, object],
    field_map: Mapping[str, str],
    projected_results: Mapping[tuple[str, int], str] | None = None,
) -> ResultValidationPreview:
    """Validate and classify one mapped row without creating database state.

    ``projected_results`` lets a caller reproduce the ordered semantics of an
    ingestion batch while keeping a preview read-only. Keys are
    ``(source_result_id, source_revision)`` and values are canonical payload
    digests. The caller owns projection order; this function never mutates the
    supplied mapping.
    """

    payload, errors = _mapped_payload(row, field_map)
    payload_digest = _sha256(_canonical_bytes(payload))
    diagnostic_payload = _jsonb_safe_projection(payload) if errors else payload
    assert isinstance(diagnostic_payload, dict)
    if errors:
        return ResultValidationPreview(
            normalized_payload=diagnostic_payload,
            payload_digest=payload_digest,
            validation_errors=errors,
            classification=ResultPreviewClassification.QUARANTINED,
        )

    source_result_id = str(payload["source_result_id"])
    source_revision = payload["source_revision"]
    assert isinstance(source_revision, int)
    projection = projected_results or {}

    if source_revision > 1:
        predecessor_key = (source_result_id, source_revision - 1)
        predecessor_exists = (
            predecessor_key in projection
            or PublishedSourceResult.objects.filter(
                organization=organization,
                source_result_id=source_result_id,
                source_revision=source_revision - 1,
            ).exists()
        )
        if not predecessor_exists:
            continuity_error = {
                "field": "source_revision",
                "code": "noncontiguous_revision",
                "message": (
                    f"Revision {source_revision} requires published revision "
                    f"{source_revision - 1} for this source result."
                ),
            }
            return ResultValidationPreview(
                normalized_payload=payload,
                payload_digest=payload_digest,
                validation_errors=[continuity_error],
                classification=ResultPreviewClassification.QUARANTINED,
            )

    current_key = (source_result_id, source_revision)
    existing_digest = projection.get(current_key)
    if existing_digest is None:
        existing_digest = (
            PublishedSourceResult.objects.filter(
                organization=organization,
                source_result_id=source_result_id,
                source_revision=source_revision,
            )
            .values_list("payload_digest", flat=True)
            .first()
        )
    if existing_digest is None:
        classification = ResultPreviewClassification.PUBLISHABLE
    elif existing_digest == payload_digest:
        classification = ResultPreviewClassification.DUPLICATE
    else:
        classification = ResultPreviewClassification.CONFLICT
    return ResultValidationPreview(
        normalized_payload=payload,
        payload_digest=payload_digest,
        validation_errors=[],
        classification=classification,
    )


def _finish_idempotency(record: IdempotencyRecord, run: IngestionRun) -> None:
    record.status = IdempotencyRecord.Status.COMPLETED
    record.response = {"run_id": str(run.run_id)}
    record.save(update_fields=["status", "response", "updated_at"])


def _outcome_from_record(record: IdempotencyRecord) -> IngestionOutcome | None:
    response = record.response
    if record.status != IdempotencyRecord.Status.COMPLETED or not isinstance(response, dict):
        return None
    run_id = response.get("run_id")
    if not isinstance(run_id, str):
        return None
    run = (
        IngestionRun.objects.select_related("artifact", "mapping_template")
        .filter(pk=run_id)
        .first()
    )
    return IngestionOutcome(run=run, replayed=True) if run is not None else None


@transaction.atomic
def _ingest_rows(
    *,
    actor: Account,
    organization: PartnerOrganization,
    mapping_template: MappingTemplate,
    operator_key: str,
    rows: Sequence[Mapping[str, object]],
    artifact_kind: SourceArtifact.Kind,
    artifact_digest: str,
    artifact_name: str,
    content_type: str,
    byte_size: int,
    object_reference: str,
    source_sheet_name: str = "",
    source_partition_key: str = "",
    source_row_numbers: Sequence[int | None] | None = None,
) -> IngestionOutcome:
    _authorize(actor, organization)
    # One tenant lock serializes artifact, run, and source-revision check/create
    # operations even when callers use different idempotency keys.
    current_organization = (
        PartnerOrganization.objects.select_for_update().filter(pk=organization.pk).first()
    )
    if (
        current_organization is None
        or current_organization.status != PartnerOrganization.Status.ACTIVE
    ):
        raise ValidationError("source organization is not active")
    current_mapping = _validate_context(
        organization=current_organization, mapping_template=mapping_template
    )
    if not operator_key.strip() or len(operator_key) > 200:
        raise ValidationError("operator key must contain 1 to 200 characters")
    if len(source_sheet_name) > 128:
        raise ValidationError("source sheet name must contain at most 128 characters")
    if len(source_partition_key) > 200:
        raise ValidationError("source partition key must contain at most 200 characters")
    if source_row_numbers is not None and len(source_row_numbers) != len(rows):
        raise ValidationError("source row coordinates do not match the intake rows")

    field_map_digest = _sha256(_canonical_bytes(current_mapping.field_map))
    request_identity = {
        "organization_id": str(current_organization.pk),
        "artifact_digest": artifact_digest,
        "artifact_kind": artifact_kind,
        "mapping_template_id": str(current_mapping.pk),
        "mapping_version": current_mapping.version,
        "field_map_digest": field_map_digest,
        "source_sheet_name": source_sheet_name,
    }
    if source_partition_key:
        request_identity["source_partition_key"] = source_partition_key
        request_identity["partition_payload_digest"] = _sha256(
            _canonical_bytes(
                {
                    "rows": list(rows),
                    "source_row_numbers": list(source_row_numbers or []),
                }
            )
        )
    request_digest = _sha256(_canonical_bytes(request_identity))
    record, _ = claim_idempotency(
        scope=f"results-ingestion:{current_organization.pk}",
        key=operator_key.strip(),
        request_digest=request_digest,
    )
    replay = _outcome_from_record(record)
    if replay is not None:
        return replay

    artifact, _ = SourceArtifact.objects.get_or_create(
        organization=current_organization,
        kind=artifact_kind,
        digest=artifact_digest,
        defaults={
            "original_name": artifact_name,
            "content_type": content_type,
            "byte_size": byte_size,
            "object_reference": object_reference,
            "uploaded_by": actor,
        },
    )
    existing_run = (
        IngestionRun.objects.select_related("artifact", "mapping_template")
        .filter(
            artifact=artifact,
            mapping_template=current_mapping,
            source_sheet_name=source_sheet_name,
            source_partition_key=source_partition_key,
        )
        .first()
    )
    if existing_run is not None:
        if existing_run.request_digest != request_digest:
            raise ValidationError("source partition was already ingested with different rows")
        _finish_idempotency(record, existing_run)
        return IngestionOutcome(run=existing_run, replayed=True)

    run = IngestionRun.objects.create(
        organization=current_organization,
        artifact=artifact,
        mapping_template=current_mapping,
        field_map_digest=field_map_digest,
        source_sheet_name=source_sheet_name,
        source_partition_key=source_partition_key,
        operator_key=operator_key.strip(),
        request_digest=request_digest,
        created_by=actor,
    )
    published_count = quarantined_count = duplicate_count = conflict_count = 0

    for row_number, row in enumerate(rows, start=1):
        source_row_number = (
            source_row_numbers[row_number - 1] if source_row_numbers is not None else None
        )
        payload, errors = _mapped_payload(row, current_mapping.field_map)
        payload_digest = _sha256(_canonical_bytes(payload))
        source_result_id = str(payload["source_result_id"])
        revision_value = payload["source_revision"]
        source_revision = revision_value if isinstance(revision_value, int) else None
        if errors:
            diagnostic_payload = _jsonb_safe_projection(payload)
            assert isinstance(diagnostic_payload, dict)
            StagedResult.objects.create(
                run=run,
                row_number=row_number,
                source_sheet_name=source_sheet_name,
                source_row_number=source_row_number,
                source_result_id=str(diagnostic_payload["source_result_id"])[
                    : StagedResult._meta.get_field("source_result_id").max_length
                ],
                source_revision=source_revision,
                normalized_payload=diagnostic_payload,
                payload_digest=payload_digest,
                validation_errors=errors,
                outcome=StagedResult.Outcome.QUARANTINED,
            )
            quarantined_count += 1
            continue

        assert source_revision is not None
        predecessor = None
        if source_revision > 1:
            predecessor = (
                PublishedSourceResult.objects.select_for_update()
                .filter(
                    organization=current_organization,
                    source_result_id=source_result_id,
                    source_revision=source_revision - 1,
                )
                .first()
            )
            if predecessor is None:
                continuity_error = {
                    "field": "source_revision",
                    "code": "noncontiguous_revision",
                    "message": (
                        f"Revision {source_revision} requires published revision "
                        f"{source_revision - 1} for this source result."
                    ),
                }
                StagedResult.objects.create(
                    run=run,
                    row_number=row_number,
                    source_sheet_name=source_sheet_name,
                    source_row_number=source_row_number,
                    source_result_id=source_result_id,
                    source_revision=source_revision,
                    normalized_payload=payload,
                    payload_digest=payload_digest,
                    validation_errors=[continuity_error],
                    outcome=StagedResult.Outcome.QUARANTINED,
                )
                quarantined_count += 1
                continue
        existing = (
            PublishedSourceResult.objects.select_for_update()
            .filter(
                organization=current_organization,
                source_result_id=source_result_id,
                source_revision=source_revision,
            )
            .first()
        )
        if existing is not None:
            outcome = (
                StagedResult.Outcome.DUPLICATE
                if existing.payload_digest == payload_digest
                else StagedResult.Outcome.CONFLICT
            )
            staged = StagedResult.objects.create(
                run=run,
                row_number=row_number,
                source_sheet_name=source_sheet_name,
                source_row_number=source_row_number,
                source_result_id=source_result_id,
                source_revision=source_revision,
                normalized_payload=payload,
                payload_digest=payload_digest,
                validation_errors=[],
                outcome=outcome,
            )
            if outcome == StagedResult.Outcome.DUPLICATE:
                duplicate_count += 1
            else:
                conflict_count += 1
                ReconciliationCase.objects.create(
                    case_type=ReconciliationCase.CaseType.SOURCE_PAYLOAD_CONFLICT,
                    staged_result=staged,
                    existing_published_result=existing,
                    details={
                        "existing_payload_digest": existing.payload_digest,
                        "submitted_payload_digest": payload_digest,
                    },
                )
            continue

        staged = StagedResult.objects.create(
            run=run,
            row_number=row_number,
            source_sheet_name=source_sheet_name,
            source_row_number=source_row_number,
            source_result_id=source_result_id,
            source_revision=source_revision,
            normalized_payload=payload,
            payload_digest=payload_digest,
            validation_errors=[],
            outcome=StagedResult.Outcome.PUBLISHED,
        )
        published = PublishedSourceResult.objects.create(
            organization=current_organization,
            source_result_id=source_result_id,
            source_revision=source_revision,
            discipline=str(payload["discipline"]),
            score_type=str(payload["score_type"]),
            normalized_payload=payload,
            payload_digest=payload_digest,
            artifact=artifact,
            staged_result=staged,
            predecessor=predecessor,
            person=None,
        )
        ReconciliationCase.objects.create(
            case_type=ReconciliationCase.CaseType.IDENTITY_UNRESOLVED,
            staged_result=staged,
            existing_published_result=published,
            details={"reason": "source identity requires explicit reconciliation"},
        )
        published_count += 1

    run.total_rows = len(rows)
    run.published_rows = published_count
    run.quarantined_rows = quarantined_count
    run.duplicate_rows = duplicate_count
    run.conflict_rows = conflict_count
    run.status = (
        IngestionRun.Status.COMPLETED_WITH_QUARANTINE
        if quarantined_count or conflict_count
        else IngestionRun.Status.COMPLETED
    )
    run.completed_at = timezone.now()
    run.save(
        update_fields=[
            "total_rows",
            "published_rows",
            "quarantined_rows",
            "duplicate_rows",
            "conflict_rows",
            "status",
            "completed_at",
        ]
    )
    _finish_idempotency(record, run)
    record_audit_event(
        actor_id=actor.account_id,
        action="results.ingestion_completed",
        target_type="ingestion_run",
        target_id=str(run.run_id),
        payload_digest=request_digest,
        metadata={
            "total_rows": run.total_rows,
            "published_rows": published_count,
            "quarantined_rows": quarantined_count,
            "duplicate_rows": duplicate_count,
            "conflict_rows": conflict_count,
        },
    )
    return IngestionOutcome(run=run, replayed=False)


def ingest_preparsed_rows(
    *,
    actor: Account,
    organization: PartnerOrganization,
    mapping_template: MappingTemplate,
    operator_key: str,
    rows: Sequence[Mapping[str, object]],
    artifact_kind: SourceArtifact.Kind,
    artifact_digest: str,
    artifact_name: str,
    content_type: str,
    byte_size: int,
    object_reference: str,
    source_sheet_name: str,
    source_partition_key: str,
    source_row_numbers: Sequence[int | None],
) -> IngestionOutcome:
    """Ingest one bounded, already-parsed source partition through ordinary intake.

    Callers remain responsible for parsing and for deriving a stable partition key.
    Validation, tenant locking, publication, identity quarantine, and idempotency all
    remain owned by the ordinary Results Desk transaction.
    """

    _authorize(actor, organization)
    if not rows:
        raise ValidationError("pre-parsed intake must contain at least one row")
    if len(rows) > MAX_ROWS:
        raise ValidationError("pre-parsed intake exceeds the safe row limit")
    if len(source_row_numbers) != len(rows):
        raise ValidationError("source row coordinates do not match the intake rows")
    if any(
        isinstance(row_number, bool) or not isinstance(row_number, int) or row_number < 1
        for row_number in source_row_numbers
    ):
        raise ValidationError("source row coordinates must be positive integers")
    if not source_partition_key.strip() or len(source_partition_key) > 200:
        raise ValidationError("source partition key must contain 1 to 200 characters")
    if not source_sheet_name or len(source_sheet_name) > 128:
        raise ValidationError("source sheet name must contain 1 to 128 characters")
    if (
        len(artifact_digest) != 64
        or artifact_digest.lower() != artifact_digest
        or any(character not in "0123456789abcdef" for character in artifact_digest)
    ):
        raise ValidationError("artifact digest must be a lowercase SHA-256 digest")
    if byte_size < 1:
        raise ValidationError("artifact byte size must be positive")
    manifest = _canonical_bytes(list(rows))
    if len(manifest) > MAX_FILE_BYTES:
        raise ValidationError("pre-parsed intake exceeds the safe batch-size limit")
    return _ingest_rows(
        actor=actor,
        organization=organization,
        mapping_template=mapping_template,
        operator_key=operator_key,
        rows=rows,
        artifact_kind=artifact_kind,
        artifact_digest=artifact_digest,
        artifact_name=artifact_name,
        content_type=content_type,
        byte_size=byte_size,
        object_reference=object_reference,
        source_sheet_name=source_sheet_name,
        source_partition_key=source_partition_key.strip(),
        source_row_numbers=source_row_numbers,
    )


def ingest_manual_rows(
    *,
    actor: Account,
    organization: PartnerOrganization,
    mapping_template: MappingTemplate,
    operator_key: str,
    rows: Sequence[Mapping[str, object]],
) -> IngestionOutcome:
    _authorize(actor, organization)
    if not rows:
        raise ValidationError("manual entry must contain at least one row")
    if len(rows) > MAX_ROWS:
        raise ValidationError("manual entry exceeds the safe row limit")
    manifest = _canonical_bytes(list(rows))
    if len(manifest) > MAX_FILE_BYTES:
        raise ValidationError("manual entry exceeds the safe batch-size limit")
    return _ingest_rows(
        actor=actor,
        organization=organization,
        mapping_template=mapping_template,
        operator_key=operator_key,
        rows=rows,
        artifact_kind=SourceArtifact.Kind.MANUAL_MANIFEST,
        artifact_digest=_sha256(manifest),
        artifact_name="manual-entry.json",
        content_type="application/vnd.mnemex.manual+json",
        byte_size=len(manifest),
        object_reference="",
    )


def ingest_spreadsheet_bytes(
    *,
    actor: Account,
    organization: PartnerOrganization,
    mapping_template: MappingTemplate,
    operator_key: str,
    filename: str,
    content_type: str,
    content: bytes,
    object_reference: str = "",
    worksheet_name: str | None = None,
) -> IngestionOutcome:
    _authorize(actor, organization)
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValidationError("spreadsheet must contain 1 byte to 5 MiB")
    safe_name = PurePath(filename).name
    if safe_name != filename or not safe_name:
        raise ValidationError("spreadsheet filename must be a plain filename")
    extension = PurePath(safe_name).suffix.lower()
    if extension == ".csv" and content_type in _CSV_TYPES:
        if (worksheet_name or "").strip():
            raise ValidationError("worksheet name can only be used with an XLSX file")
        parsed = _parse_csv(content)
    elif extension == ".xlsx" and content_type in _XLSX_TYPES:
        parsed = _parse_xlsx(content, worksheet_name=worksheet_name)
    else:
        raise ValidationError("only matching CSV and XLSX file types are accepted")
    if not parsed.rows:
        raise ValidationError("spreadsheet must contain at least one data row")
    return _ingest_rows(
        actor=actor,
        organization=organization,
        mapping_template=mapping_template,
        operator_key=operator_key,
        rows=parsed.rows,
        artifact_kind=SourceArtifact.Kind.SPREADSHEET,
        artifact_digest=_sha256(content),
        artifact_name=safe_name,
        content_type=content_type,
        byte_size=len(content),
        object_reference=object_reference,
        source_sheet_name=parsed.source_sheet_name,
        source_row_numbers=parsed.source_row_numbers,
    )
