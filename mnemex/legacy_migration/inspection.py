from __future__ import annotations

import hashlib
import json
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath

from django.core.exceptions import ValidationError
from openpyxl import load_workbook  # type: ignore[import-untyped]

MAX_LEGACY_FILE_BYTES = 20 * 1024 * 1024
MAX_LEGACY_ARCHIVE_ENTRIES = 512
MAX_LEGACY_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_LEGACY_SHEETS = 64
MAX_LEGACY_ROWS_PER_SHEET = 25_000
MAX_LEGACY_COLUMNS_PER_SHEET = 256
MAX_LEGACY_CELL_CHARACTERS = 10_000
MAX_LEGACY_CELL_WORK_PER_SHEET = 250_000
MAX_LEGACY_CELL_WORK_PER_WORKBOOK = 1_000_000
MAX_HEADER_CANDIDATE_ROW = 25


@dataclass(frozen=True)
class LegacyWorksheetInspection:
    index: int
    name: str
    visibility: str
    max_row: int
    max_column: int
    merged_ranges: tuple[str, ...]
    header_candidates: tuple[int, ...]
    formula_cells: tuple[str, ...]
    metadata_digest: str


@dataclass(frozen=True)
class LegacyWorkbookInspection:
    artifact_digest: str
    source_state_digest: str
    archive_entries: int
    uncompressed_bytes: int
    sheets: tuple[LegacyWorksheetInspection, ...]


def _canonical_bytes(value: object) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _visibility(sheet_state: str) -> str:
    if sheet_state == "veryHidden":
        return "very_hidden"
    if sheet_state in {"visible", "hidden"}:
        return sheet_state
    raise ValidationError("XLSX worksheet visibility is unsupported")


def _validate_archive(content: bytes) -> tuple[int, int]:
    if not content or len(content) > MAX_LEGACY_FILE_BYTES:
        raise ValidationError("legacy XLSX must contain 1 byte to 20 MiB")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_LEGACY_ARCHIVE_ENTRIES:
                raise ValidationError("legacy XLSX archive exceeds the safe entry limit")
            names: set[str] = set()
            total_size = 0
            for entry in entries:
                normalized_name = entry.filename.replace("\\", "/")
                path = PurePosixPath(normalized_name)
                if (
                    not normalized_name
                    or normalized_name.startswith("/")
                    or ".." in path.parts
                    or entry.flag_bits & 0x1
                ):
                    raise ValidationError("legacy XLSX archive contains an unsafe entry")
                if normalized_name in names:
                    raise ValidationError("legacy XLSX archive contains duplicate entries")
                names.add(normalized_name)
                total_size += entry.file_size
                if total_size > MAX_LEGACY_UNCOMPRESSED_BYTES:
                    raise ValidationError(
                        "legacy XLSX archive exceeds the safe uncompressed-size limit"
                    )
            return len(entries), total_size
    except zipfile.BadZipFile as exc:
        raise ValidationError("legacy XLSX is not a valid workbook archive") from exc


def _header_candidates(sheet: object, *, max_row: int, max_column: int) -> tuple[int, ...]:
    candidates: list[int] = []
    candidate_limit = min(max_row, MAX_HEADER_CANDIDATE_ROW)
    for row_number in range(1, candidate_limit + 1):
        values: list[str] = []
        for column_number in range(1, max_column + 1):
            cell = sheet.cell(row=row_number, column=column_number)  # type: ignore[attr-defined]
            if cell.value is None:
                continue
            text = str(cell.value).strip()
            if len(text) > MAX_LEGACY_CELL_CHARACTERS:
                raise ValidationError("legacy XLSX cell exceeds the safe character limit")
            if text:
                values.append(text)
        if len(values) >= 2 and len(set(values)) == len(values):
            candidates.append(row_number)
    return tuple(candidates)


def inspect_legacy_workbook(content: bytes) -> LegacyWorkbookInspection:
    """Return deterministic workbook structure without exposing row values."""

    archive_entries, uncompressed_bytes = _validate_archive(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise ValidationError("legacy XLSX workbook could not be parsed safely") from exc
    try:
        if len(workbook.worksheets) > MAX_LEGACY_SHEETS:
            raise ValidationError("legacy XLSX exceeds the safe sheet limit")
        inspected_sheets: list[LegacyWorksheetInspection] = []
        workbook_cell_work = 0
        for index, sheet in enumerate(workbook.worksheets):
            max_row = int(sheet.max_row or 0)
            max_column = int(sheet.max_column or 0)
            if max_row > MAX_LEGACY_ROWS_PER_SHEET:
                raise ValidationError("legacy XLSX exceeds the safe row limit")
            if max_column > MAX_LEGACY_COLUMNS_PER_SHEET:
                raise ValidationError("legacy XLSX exceeds the safe column limit")
            sheet_cell_work = max_row * max_column
            workbook_cell_work += sheet_cell_work
            if (
                sheet_cell_work > MAX_LEGACY_CELL_WORK_PER_SHEET
                or workbook_cell_work > MAX_LEGACY_CELL_WORK_PER_WORKBOOK
            ):
                raise ValidationError("legacy XLSX exceeds the safe cell-work limit")
            candidates = _header_candidates(
                sheet,
                max_row=max_row,
                max_column=max_column,
            )
            formulas: list[str] = []
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and len(value) > MAX_LEGACY_CELL_CHARACTERS:
                        raise ValidationError("legacy XLSX cell exceeds the safe character limit")
                    if cell.data_type == "f":
                        formulas.append(cell.coordinate)
            visibility = _visibility(sheet.sheet_state)
            merged_ranges = tuple(str(item) for item in sheet.merged_cells.ranges)
            metadata = {
                "index": index,
                "name": sheet.title,
                "visibility": visibility,
                "max_row": max_row,
                "max_column": max_column,
                "merged_ranges": merged_ranges,
                "header_candidates": candidates,
                "formula_cells": formulas,
            }
            inspected_sheets.append(
                LegacyWorksheetInspection(
                    index=index,
                    name=sheet.title,
                    visibility=visibility,
                    max_row=max_row,
                    max_column=max_column,
                    merged_ranges=merged_ranges,
                    header_candidates=candidates,
                    formula_cells=tuple(formulas),
                    metadata_digest=_sha256(_canonical_bytes(metadata)),
                )
            )
        artifact_digest = _sha256(content)
        source_state = {
            "artifact_digest": artifact_digest,
            "archive_entries": archive_entries,
            "uncompressed_bytes": uncompressed_bytes,
            "sheets": [
                {
                    "index": sheet.index,
                    "metadata_digest": sheet.metadata_digest,
                }
                for sheet in inspected_sheets
            ],
        }
        return LegacyWorkbookInspection(
            artifact_digest=artifact_digest,
            source_state_digest=_sha256(_canonical_bytes(source_state)),
            archive_entries=archive_entries,
            uncompressed_bytes=uncompressed_bytes,
            sheets=tuple(inspected_sheets),
        )
    finally:
        workbook.close()
