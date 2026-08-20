from __future__ import annotations

import hashlib
import json
import zipfile
from collections.abc import Collection, Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from uuid import UUID

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import models, transaction
from django.utils import timezone
from openpyxl import load_workbook  # type: ignore[import-untyped]

from mnemex.accounts.authorization import Action, may_perform
from mnemex.accounts.models import Account
from mnemex.legacy_migration.inspection import (
    MAX_LEGACY_COLUMNS_PER_SHEET,
    MAX_LEGACY_ROWS_PER_SHEET,
    LegacyWorkbookInspection,
    LegacyWorksheetInspection,
    inspect_legacy_workbook,
)
from mnemex.legacy_migration.models import (
    MAX_LEGACY_APPLY_CHUNK_ROWS,
    LegacyMigrationCheckpoint,
    LegacyMigrationDecisionRevision,
    LegacyMigrationJob,
    LegacyMigrationManifest,
    LegacyMigrationRowPreview,
    LegacyMigrationRun,
    LegacyTablePlan,
    LegacyWorksheetPlan,
)
from mnemex.partners.models import PartnerOrganization
from mnemex.results.models import (
    IngestionRun,
    MappingTemplate,
    PublishedSourceResult,
    SourceArtifact,
    StagedResult,
)
from mnemex.results.services import (
    ResultPreviewClassification,
    ingest_preparsed_rows,
    preview_mapped_result,
)

_REQUIRED_CANONICAL_FIELDS = (
    "source_result_id",
    "source_revision",
    "source_event_id",
    "event_name",
    "result_date",
    "competitor_name",
    "discipline",
    "score_type",
    "score",
)
_OPTIONAL_CANONICAL_FIELDS = (
    "heat_id",
    "wood_species",
    "wood_diameter_mm",
    "wood_quality",
)
_CANONICAL_FIELDS = frozenset(_REQUIRED_CANONICAL_FIELDS + _OPTIONAL_CANONICAL_FIELDS)
_SHEET_CONFIGURATION_KEYS = frozenset({"sheet_name", "disposition", "ignore_reason", "tables"})
_TABLE_CONFIGURATION_KEYS = frozenset(
    {
        "label",
        "header_row",
        "start_row",
        "end_row",
        "start_column",
        "end_column",
        "mapping_rules",
    }
)
_RULE_KINDS = frozenset({"column", "constant", "derived"})
_MAX_TABLE_ROWS = 5_000
_MAX_TABLES_PER_SHEET = 32
_MAX_TABLES_PER_RUN = 128
_MAX_PLANNED_ROWS_PER_RUN = 25_000
MAX_APPLY_CHUNK_ROWS = MAX_LEGACY_APPLY_CHUNK_ROWS

_CANONICAL_RESULT_FIELDS = (
    "source_result_id",
    "source_revision",
    "source_event_id",
    "event_name",
    "result_date",
    "competitor_name",
    "discipline",
    "score_type",
    "score",
    "heat_id",
    "wood_species",
    "wood_diameter_mm",
    "wood_quality",
)


@dataclass(frozen=True)
class _ValidatedTableConfiguration:
    label: str
    header_row: int
    start_row: int
    end_row: int
    start_column: int
    end_column: int
    mapping_rules: dict[str, dict[str, object]]
    mapping_digest: str


@dataclass(frozen=True)
class _ValidatedSheetConfiguration:
    inspection: LegacyWorksheetInspection
    disposition: str
    ignore_reason: str
    tables: tuple[_ValidatedTableConfiguration, ...]


@dataclass(frozen=True)
class _PreviewFact:
    table_plan: LegacyTablePlan
    source_sheet_index: int
    source_row_number: int
    source_column_start: int
    source_column_end: int
    canonical_payload: dict[str, object]
    payload_digest: str
    classification: str
    validation_errors: list[dict[str, str]]


@dataclass(frozen=True)
class _ApplyChunk:
    sequence: int
    table_plan: LegacyTablePlan
    previews: tuple[LegacyMigrationRowPreview, ...]
    digest: str


class ApplyClaimLost(PermissionDenied):
    """The caller no longer owns the fenced APPLY lease."""


def _canonical_bytes(value: object) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")


def _sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _authorize(actor: Account, run: LegacyMigrationRun) -> None:
    if not may_perform(actor, Action.MANAGE_RESULTS, organization=run.organization):
        raise PermissionDenied("an MFA-bound results manager role is required")


def _current_run(run: LegacyMigrationRun) -> LegacyMigrationRun:
    current = (
        LegacyMigrationRun.objects.select_for_update()
        .select_related("organization", "inventory__artifact")
        .filter(pk=run.pk)
        .first()
    )
    if current is None:
        raise ValidationError("legacy migration run no longer exists")
    return current


def _verify_source(run: LegacyMigrationRun, workbook_content: bytes) -> LegacyWorkbookInspection:
    inspection = inspect_legacy_workbook(workbook_content)
    if inspection.artifact_digest != run.inventory.artifact.digest:
        raise ValidationError("workbook content does not match the source artifact digest")
    return inspection


def _plain_string(value: object, *, field: str, maximum: int) -> str:
    if not isinstance(value, str) or not value.strip() or len(value.strip()) > maximum:
        raise ValidationError(f"{field} must contain 1 to {maximum} characters")
    return value.strip()


def _positive_integer(value: object, *, field: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise ValidationError(f"{field} must be a positive integer")
    return value


def _validate_mapping_rules(value: object) -> tuple[dict[str, dict[str, object]], str]:
    if not isinstance(value, Mapping):
        raise ValidationError("mapping rules must be an object")
    unknown_fields = sorted(set(value) - _CANONICAL_FIELDS)
    if unknown_fields:
        raise ValidationError(
            f"mapping rules contain unknown canonical fields: {', '.join(unknown_fields)}"
        )
    missing_fields = [field for field in _REQUIRED_CANONICAL_FIELDS if field not in value]
    if missing_fields:
        raise ValidationError(
            f"mapping rules are missing canonical fields: {', '.join(missing_fields)}"
        )
    normalized: dict[str, dict[str, object]] = {}
    for canonical_field, raw_rule in value.items():
        if not isinstance(canonical_field, str) or not isinstance(raw_rule, Mapping):
            raise ValidationError("each mapping rule must be an object")
        kind = raw_rule.get("kind")
        if kind not in _RULE_KINDS:
            raise ValidationError(f"unknown mapping rule kind for {canonical_field}")
        if kind == "column":
            if set(raw_rule) != {"kind", "column"}:
                raise ValidationError("column mapping rules only accept kind and column")
            column = _plain_string(
                raw_rule.get("column"),
                field="mapping column",
                maximum=255,
            )
            normalized[canonical_field] = {"kind": "column", "column": column}
        elif kind == "constant":
            if set(raw_rule) != {"kind", "value"}:
                raise ValidationError("constant mapping rules only accept kind and value")
            constant = raw_rule.get("value")
            if constant is not None and not isinstance(constant, (str, int, float, bool)):
                raise ValidationError("constant mapping values must be scalar JSON values")
            normalized[canonical_field] = {"kind": "constant", "value": constant}
        else:
            if set(raw_rule) != {"kind", "name"}:
                raise ValidationError("derived mapping rules only accept kind and name")
            if raw_rule.get("name") != "artifact_coordinate_id":
                raise ValidationError("unknown derived mapping rule")
            if canonical_field != "source_result_id":
                raise ValidationError("artifact_coordinate_id may only populate source_result_id")
            normalized[canonical_field] = {
                "kind": "derived",
                "name": "artifact_coordinate_id",
            }
    digest = _sha256(_canonical_bytes(normalized))
    return normalized, digest


def _validate_table_configuration(value: object) -> _ValidatedTableConfiguration:
    if not isinstance(value, Mapping):
        raise ValidationError("table configurations must be objects")
    unknown_keys = sorted(set(value) - _TABLE_CONFIGURATION_KEYS)
    if unknown_keys:
        raise ValidationError(
            f"table configuration contains unknown keys: {', '.join(unknown_keys)}"
        )
    label = _plain_string(value.get("label"), field="table label", maximum=160)
    header_row = _positive_integer(value.get("header_row"), field="header row")
    start_row = _positive_integer(value.get("start_row"), field="start row")
    end_row = _positive_integer(value.get("end_row"), field="end row")
    start_column = _positive_integer(value.get("start_column"), field="start column")
    end_column = _positive_integer(value.get("end_column"), field="end column")
    if end_row < start_row or not start_row <= header_row <= end_row:
        raise ValidationError("table header and row bounds are invalid")
    if end_column < start_column:
        raise ValidationError("table column bounds are invalid")
    if end_row - header_row > _MAX_TABLE_ROWS:
        raise ValidationError("legacy table exceeds the safe row limit")
    if end_row > MAX_LEGACY_ROWS_PER_SHEET:
        raise ValidationError("legacy table exceeds the safe row limit")
    if end_column > MAX_LEGACY_COLUMNS_PER_SHEET:
        raise ValidationError("legacy table exceeds the safe column limit")
    rules, mapping_digest = _validate_mapping_rules(value.get("mapping_rules"))
    return _ValidatedTableConfiguration(
        label=label,
        header_row=header_row,
        start_row=start_row,
        end_row=end_row,
        start_column=start_column,
        end_column=end_column,
        mapping_rules=rules,
        mapping_digest=mapping_digest,
    )


def _table_regions_overlap(
    first: _ValidatedTableConfiguration,
    second: _ValidatedTableConfiguration,
) -> bool:
    rows_overlap = first.start_row <= second.end_row and second.start_row <= first.end_row
    columns_overlap = (
        first.start_column <= second.end_column and second.start_column <= first.end_column
    )
    return rows_overlap and columns_overlap


def _load_workbook(content: bytes) -> object:
    try:
        return load_workbook(
            BytesIO(content),
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise ValidationError("legacy XLSX workbook could not be parsed safely") from exc


def _headers_for_table(sheet: object, table: _ValidatedTableConfiguration) -> dict[str, int]:
    headers: list[str] = []
    for column_number in range(table.start_column, table.end_column + 1):
        cell = sheet.cell(row=table.header_row, column=column_number)  # type: ignore[attr-defined]
        if cell.data_type == "f":
            raise ValidationError(
                "formula-backed mapped cells require explicit review before migration"
            )
        header = str(cell.value).strip() if cell.value is not None else ""
        headers.append(header)
    if any(not header for header in headers):
        raise ValidationError("legacy table header names cannot be blank")
    if len(set(headers)) != len(headers):
        raise ValidationError("legacy table header names must be unique")
    header_columns = {header: table.start_column + offset for offset, header in enumerate(headers)}
    for rule in table.mapping_rules.values():
        if rule["kind"] == "column" and rule["column"] not in header_columns:
            raise ValidationError(f"mapped column {rule['column']!r} was not found in the table")
    return header_columns


def _validate_sheet_configurations(
    *,
    inspection: LegacyWorkbookInspection,
    workbook_content: bytes,
    sheet_configurations: Sequence[Mapping[str, object]],
) -> tuple[_ValidatedSheetConfiguration, ...]:
    if not isinstance(sheet_configurations, Sequence) or isinstance(
        sheet_configurations, (str, bytes)
    ):
        raise ValidationError("sheet configurations must be a list")
    by_name: dict[str, Mapping[str, object]] = {}
    preliminary: dict[str, tuple[str, str, tuple[_ValidatedTableConfiguration, ...]]] = {}
    total_tables = 0
    total_planned_rows = 0
    for raw_configuration in sheet_configurations:
        if not isinstance(raw_configuration, Mapping):
            raise ValidationError("sheet configurations must be objects")
        unknown_keys = sorted(set(raw_configuration) - _SHEET_CONFIGURATION_KEYS)
        if unknown_keys:
            raise ValidationError(
                f"sheet configuration contains unknown keys: {', '.join(unknown_keys)}"
            )
        sheet_name = raw_configuration.get("sheet_name")
        if not isinstance(sheet_name, str) or not sheet_name:
            raise ValidationError("sheet name must be an exact non-empty string")
        if sheet_name in by_name:
            raise ValidationError("each workbook sheet must be configured exactly once")
        by_name[sheet_name] = raw_configuration
        disposition = raw_configuration.get("disposition")
        raw_tables = raw_configuration.get("tables", [])
        if not isinstance(raw_tables, Sequence) or isinstance(raw_tables, (str, bytes)):
            raise ValidationError("table configurations must be a list")
        tables = tuple(_validate_table_configuration(table) for table in raw_tables)
        if len(tables) > _MAX_TABLES_PER_SHEET:
            raise ValidationError("worksheet exceeds the safe table-region limit")
        total_tables += len(tables)
        if total_tables > _MAX_TABLES_PER_RUN:
            raise ValidationError("migration exceeds the safe table-region limit")
        total_planned_rows += sum(table.end_row - table.header_row for table in tables)
        if total_planned_rows > _MAX_PLANNED_ROWS_PER_RUN:
            raise ValidationError("migration exceeds the safe planned-row limit")
        if disposition == LegacyWorksheetPlan.Disposition.INCLUDED:
            if not tables:
                raise ValidationError("included worksheets require at least one table")
            ignore_reason = ""
        elif disposition == LegacyWorksheetPlan.Disposition.IGNORED:
            if tables:
                raise ValidationError("ignored worksheets cannot contain table plans")
            ignore_reason = _plain_string(
                raw_configuration.get("ignore_reason"),
                field="ignore reason",
                maximum=240,
            )
        else:
            raise ValidationError("sheet disposition must be included or ignored")
        labels = [table.label for table in tables]
        if len(labels) != len(set(labels)):
            raise ValidationError("table labels must be unique within a worksheet")
        for index, table in enumerate(tables):
            if any(_table_regions_overlap(table, other) for other in tables[index + 1 :]):
                raise ValidationError("table regions cannot overlap within a worksheet")
        preliminary[sheet_name] = (str(disposition), ignore_reason, tables)

    expected_names = [sheet.name for sheet in inspection.sheets]
    if set(by_name) != set(expected_names) or len(by_name) != len(expected_names):
        raise ValidationError("every workbook sheet must be configured exactly once")

    workbook = _load_workbook(workbook_content)
    try:
        validated: list[_ValidatedSheetConfiguration] = []
        for inspected_sheet in inspection.sheets:
            disposition, ignore_reason, tables = preliminary[inspected_sheet.name]
            sheet = workbook.worksheets[inspected_sheet.index]  # type: ignore[attr-defined]
            for table in tables:
                _headers_for_table(sheet, table)
            validated.append(
                _ValidatedSheetConfiguration(
                    inspection=inspected_sheet,
                    disposition=disposition,
                    ignore_reason=ignore_reason,
                    tables=tables,
                )
            )
        return tuple(validated)
    finally:
        workbook.close()  # type: ignore[attr-defined]


def stable_artifact_coordinate_id(
    *,
    artifact_digest: str,
    sheet_index: int,
    sheet_name: str,
    table_digest: str,
    source_row_number: int,
) -> str:
    """Derive a stable source ID without using a competitor name."""

    coordinate = {
        "version": "legacy-coordinate-v1",
        "artifact_digest": artifact_digest,
        "sheet_index": sheet_index,
        "sheet_name": sheet_name,
        "table_digest": table_digest,
        "source_row_number": source_row_number,
    }
    return f"legacy:{_sha256(_canonical_bytes(coordinate))}"


def _table_coordinate_digest(table: LegacyTablePlan) -> str:
    """Bind a derived source ID to one exact rectangular table plan."""

    return _sha256(
        _canonical_bytes(
            {
                "version": "legacy-table-coordinate-v1",
                "label": table.label,
                "header_row": table.header_row,
                "start_row": table.start_row,
                "end_row": table.end_row,
                "start_column": table.start_column,
                "end_column": table.end_column,
                "mapping_digest": table.mapping_digest,
            }
        )
    )


@transaction.atomic
def configure_migration_run(
    *,
    actor: Account,
    run: LegacyMigrationRun,
    workbook_content: bytes,
    sheet_configurations: Sequence[Mapping[str, object]],
) -> LegacyMigrationRun:
    """Persist an explicit, immutable sheet/table plan after full validation."""

    current = _current_run(run)
    _authorize(actor, current)
    if current.state != LegacyMigrationRun.State.DISCOVERED:
        raise ValidationError("only a discovered legacy migration run can be configured")
    if current.worksheets.exists() or current.table_plans.exists():
        raise ValidationError("legacy migration run already contains plan evidence")
    inspection = _verify_source(current, workbook_content)
    configurations = _validate_sheet_configurations(
        inspection=inspection,
        workbook_content=workbook_content,
        sheet_configurations=sheet_configurations,
    )
    plan_payload = {
        "version": "legacy-plan-v1",
        "source_state_digest": inspection.source_state_digest,
        "sheets": [
            {
                "index": item.inspection.index,
                "name": item.inspection.name,
                "metadata_digest": item.inspection.metadata_digest,
                "disposition": item.disposition,
                "ignore_reason": item.ignore_reason,
                "tables": [
                    {
                        "label": table.label,
                        "header_row": table.header_row,
                        "start_row": table.start_row,
                        "end_row": table.end_row,
                        "start_column": table.start_column,
                        "end_column": table.end_column,
                        "mapping_digest": table.mapping_digest,
                    }
                    for table in item.tables
                ],
            }
            for item in configurations
        ],
    }
    plan_digest = _sha256(_canonical_bytes(plan_payload))

    included_count = 0
    ignored_count = 0
    for item in configurations:
        worksheet = LegacyWorksheetPlan.objects.create(
            organization=current.organization,
            run=current,
            sheet_index=item.inspection.index,
            sheet_name=item.inspection.name,
            visibility=item.inspection.visibility,
            max_row=item.inspection.max_row,
            max_column=item.inspection.max_column,
            disposition=item.disposition,
            ignore_reason=item.ignore_reason,
            metadata_digest=item.inspection.metadata_digest,
            header_candidates=list(item.inspection.header_candidates),
        )
        if item.disposition == LegacyWorksheetPlan.Disposition.INCLUDED:
            included_count += 1
        else:
            ignored_count += 1
        for table in item.tables:
            LegacyTablePlan.objects.create(
                organization=current.organization,
                run=current,
                worksheet=worksheet,
                label=table.label,
                header_row=table.header_row,
                start_row=table.start_row,
                end_row=table.end_row,
                start_column=table.start_column,
                end_column=table.end_column,
                mapping_rules=table.mapping_rules,
                mapping_digest=table.mapping_digest,
                created_by=actor,
            )

    current.state = LegacyMigrationRun.State.CONFIGURED
    current.plan_digest = plan_digest
    current.source_state_digest = inspection.source_state_digest
    current.total_sheets = len(configurations)
    current.included_sheets = included_count
    current.ignored_sheets = ignored_count
    current.save(
        update_fields=[
            "state",
            "plan_digest",
            "source_state_digest",
            "total_sheets",
            "included_sheets",
            "ignored_sheets",
            "updated_at",
        ]
    )
    return current


def _is_blank_row(sheet: object, table: LegacyTablePlan, source_row_number: int) -> bool:
    for column_number in range(table.start_column, table.end_column + 1):
        value = sheet.cell(  # type: ignore[attr-defined]
            row=source_row_number,
            column=column_number,
        ).value
        if value is not None and (not isinstance(value, str) or value.strip()):
            return False
    return True


def _map_row(
    *,
    sheet: object,
    worksheet: LegacyWorksheetPlan,
    table: LegacyTablePlan,
    source_row_number: int,
    artifact_digest: str,
) -> dict[str, object]:
    table_configuration = _ValidatedTableConfiguration(
        label=table.label,
        header_row=table.header_row,
        start_row=table.start_row,
        end_row=table.end_row,
        start_column=table.start_column,
        end_column=table.end_column,
        mapping_rules=table.mapping_rules,
        mapping_digest=table.mapping_digest,
    )
    headers = _headers_for_table(sheet, table_configuration)
    mapped: dict[str, object] = {}
    for canonical_field, rule in table.mapping_rules.items():
        if rule["kind"] == "column":
            column_number = headers[str(rule["column"])]
            cell = sheet.cell(  # type: ignore[attr-defined]
                row=source_row_number,
                column=column_number,
            )
            if cell.data_type == "f":
                raise ValidationError(
                    "formula-backed mapped cells require explicit review before migration"
                )
            mapped[canonical_field] = cell.value
        elif rule["kind"] == "constant":
            mapped[canonical_field] = rule.get("value")
        elif rule["kind"] == "derived":
            mapped[canonical_field] = stable_artifact_coordinate_id(
                artifact_digest=artifact_digest,
                sheet_index=worksheet.sheet_index,
                sheet_name=worksheet.sheet_name,
                table_digest=_table_coordinate_digest(table),
                source_row_number=source_row_number,
            )
        else:
            raise ValidationError(f"unknown mapping rule kind for {canonical_field}")
    return mapped


def _verify_persisted_plan(
    run: LegacyMigrationRun, inspection: LegacyWorkbookInspection
) -> list[LegacyWorksheetPlan]:
    if inspection.source_state_digest != run.source_state_digest:
        raise ValidationError("workbook source state no longer matches the configured plan")
    worksheets = list(run.worksheets.prefetch_related("table_plans").order_by("sheet_index"))
    if len(worksheets) != len(inspection.sheets):
        raise ValidationError("configured worksheet evidence is incomplete")
    for worksheet, inspected in zip(worksheets, inspection.sheets, strict=True):
        if (
            worksheet.sheet_index != inspected.index
            or worksheet.sheet_name != inspected.name
            or worksheet.visibility != inspected.visibility
            or worksheet.metadata_digest != inspected.metadata_digest
        ):
            raise ValidationError("configured worksheet evidence does not match the workbook")
        if (
            worksheet.disposition == LegacyWorksheetPlan.Disposition.INCLUDED
            and not worksheet.table_plans.exists()
        ):
            raise ValidationError("included worksheet has no configured table regions")
    return worksheets


def _manifest_payload(
    *, run: LegacyMigrationRun, preview_facts: Sequence[_PreviewFact]
) -> dict[str, object]:
    return {
        "version": "legacy-dry-run-manifest-v1",
        "run_id": str(run.pk),
        "plan_digest": run.plan_digest,
        "source_state_digest": run.source_state_digest,
        "rows": [
            {
                "sheet_index": fact.source_sheet_index,
                "table_mapping_digest": fact.table_plan.mapping_digest,
                "source_row_number": fact.source_row_number,
                "source_column_start": fact.source_column_start,
                "source_column_end": fact.source_column_end,
                "payload_digest": fact.payload_digest,
                "classification": fact.classification,
                "error_codes": [error.get("code", "") for error in fact.validation_errors],
            }
            for fact in preview_facts
        ],
    }


def _persisted_preview_manifest_payload(run: LegacyMigrationRun) -> dict[str, object]:
    rows: list[dict[str, object]] = []
    previews = run.row_previews.select_related("table_plan").order_by(
        "source_sheet_index",
        "table_plan__start_row",
        "table_plan__start_column",
        "table_plan__label",
        "table_plan_id",
        "source_row_number",
        "preview_id",
    )
    for preview in previews:
        if preview.payload_digest != _sha256(_canonical_bytes(preview.canonical_payload)):
            raise ValidationError("persisted preview evidence payload digest is invalid")
        errors = preview.validation_errors
        if not isinstance(errors, list) or not all(isinstance(error, dict) for error in errors):
            raise ValidationError("persisted preview evidence validation errors are invalid")
        rows.append(
            {
                "sheet_index": preview.source_sheet_index,
                "table_mapping_digest": preview.table_plan.mapping_digest,
                "source_row_number": preview.source_row_number,
                "source_column_start": preview.source_column_start,
                "source_column_end": preview.source_column_end,
                "payload_digest": preview.payload_digest,
                "classification": preview.classification,
                "error_codes": [str(error.get("code", "")) for error in errors],
            }
        )
    return {
        "version": "legacy-dry-run-manifest-v1",
        "run_id": str(run.pk),
        "plan_digest": run.plan_digest,
        "source_state_digest": run.source_state_digest,
        "rows": rows,
    }


@transaction.atomic
def dry_run_migration(
    *,
    actor: Account,
    run: LegacyMigrationRun,
    workbook_content: bytes,
) -> LegacyMigrationRun:
    """Persist all preview evidence atomically after a completely read-only pass."""

    current = _current_run(run)
    _authorize(actor, current)
    if current.state != LegacyMigrationRun.State.CONFIGURED:
        raise ValidationError("only a configured legacy migration run can be previewed")
    if (
        current.row_previews.exists()
        or current.manifests.filter(kind=LegacyMigrationManifest.Kind.DRY_RUN).exists()
    ):
        raise ValidationError("legacy migration run already contains dry-run evidence")
    inspection = _verify_source(current, workbook_content)
    worksheets = _verify_persisted_plan(current, inspection)
    workbook = _load_workbook(workbook_content)
    preview_facts: list[_PreviewFact] = []
    projected_results: dict[tuple[str, int], str] = {}
    try:
        for worksheet in worksheets:
            if worksheet.disposition == LegacyWorksheetPlan.Disposition.IGNORED:
                continue
            sheet = workbook.worksheets[worksheet.sheet_index]  # type: ignore[attr-defined]
            for table in worksheet.table_plans.order_by(
                "start_row", "start_column", "label", "table_plan_id"
            ):
                rules, mapping_digest = _validate_mapping_rules(table.mapping_rules)
                if mapping_digest != table.mapping_digest:
                    raise ValidationError("stored table mapping digest is invalid")
                table.mapping_rules = rules
                table_configuration = _ValidatedTableConfiguration(
                    label=table.label,
                    header_row=table.header_row,
                    start_row=table.start_row,
                    end_row=table.end_row,
                    start_column=table.start_column,
                    end_column=table.end_column,
                    mapping_rules=rules,
                    mapping_digest=mapping_digest,
                )
                _headers_for_table(sheet, table_configuration)
                identity_field_map = {field: field for field in rules}
                for source_row_number in range(table.header_row + 1, table.end_row + 1):
                    if _is_blank_row(sheet, table, source_row_number):
                        canonical_payload: dict[str, object] = {"status": "ignored_blank_row"}
                        payload_digest = _sha256(_canonical_bytes(canonical_payload))
                        classification = LegacyMigrationRowPreview.Classification.IGNORED.value
                        validation_errors: list[dict[str, str]] = []
                    else:
                        mapped = _map_row(
                            sheet=sheet,
                            worksheet=worksheet,
                            table=table,
                            source_row_number=source_row_number,
                            artifact_digest=inspection.artifact_digest,
                        )
                        result_preview = preview_mapped_result(
                            organization=current.organization,
                            row=mapped,
                            field_map=identity_field_map,
                            projected_results=projected_results,
                        )
                        canonical_payload = result_preview.normalized_payload
                        payload_digest = result_preview.payload_digest
                        classification = result_preview.classification.value
                        validation_errors = result_preview.validation_errors
                        if result_preview.classification == ResultPreviewClassification.PUBLISHABLE:
                            source_result_id = str(canonical_payload["source_result_id"])
                            source_revision = canonical_payload["source_revision"]
                            assert isinstance(source_revision, int)
                            projected_results[(source_result_id, source_revision)] = payload_digest
                    preview_facts.append(
                        _PreviewFact(
                            table_plan=table,
                            source_sheet_index=worksheet.sheet_index,
                            source_row_number=source_row_number,
                            source_column_start=table.start_column,
                            source_column_end=table.end_column,
                            canonical_payload=canonical_payload,
                            payload_digest=payload_digest,
                            classification=str(classification),
                            validation_errors=validation_errors,
                        )
                    )
    finally:
        workbook.close()  # type: ignore[attr-defined]

    manifest_payload = _manifest_payload(run=current, preview_facts=preview_facts)
    manifest_digest = _sha256(_canonical_bytes(manifest_payload))
    for fact in preview_facts:
        LegacyMigrationRowPreview.objects.create(
            organization=current.organization,
            run=current,
            table_plan=fact.table_plan,
            source_sheet_index=fact.source_sheet_index,
            source_row_number=fact.source_row_number,
            source_column_start=fact.source_column_start,
            source_column_end=fact.source_column_end,
            canonical_payload=fact.canonical_payload,
            payload_digest=fact.payload_digest,
            classification=fact.classification,
            validation_errors=fact.validation_errors,
            source_state_digest=current.source_state_digest,
        )
    LegacyMigrationManifest.objects.create(
        organization=current.organization,
        run=current,
        kind=LegacyMigrationManifest.Kind.DRY_RUN,
        sequence=1,
        digest=manifest_digest,
        payload=manifest_payload,
        created_by=actor,
    )

    counts = {
        classification: sum(fact.classification == classification for fact in preview_facts)
        for classification in LegacyMigrationRowPreview.Classification.values
    }
    current.state = LegacyMigrationRun.State.DRY_RUN_COMPLETE
    current.dry_run_manifest_digest = manifest_digest
    current.total_rows = len(preview_facts)
    current.ignored_rows = counts[LegacyMigrationRowPreview.Classification.IGNORED]
    current.candidate_rows = current.total_rows - current.ignored_rows
    current.publishable_rows = counts[LegacyMigrationRowPreview.Classification.PUBLISHABLE]
    current.quarantined_rows = counts[LegacyMigrationRowPreview.Classification.QUARANTINED]
    current.duplicate_rows = counts[LegacyMigrationRowPreview.Classification.DUPLICATE]
    current.conflict_rows = counts[LegacyMigrationRowPreview.Classification.CONFLICT]
    current.save(
        update_fields=[
            "state",
            "dry_run_manifest_digest",
            "total_rows",
            "candidate_rows",
            "ignored_rows",
            "publishable_rows",
            "quarantined_rows",
            "duplicate_rows",
            "conflict_rows",
            "updated_at",
        ]
    )
    return current


def _authorize_reviewer(actor: Account, run: LegacyMigrationRun) -> None:
    if not may_perform(actor, Action.REVIEW_EXPORT, organization=run.organization):
        raise PermissionDenied("an MFA-bound export reviewer role is required")


def _rationale(value: str) -> str:
    return _plain_string(value, field="decision rationale", maximum=500)


def _lock_organization(run: LegacyMigrationRun) -> PartnerOrganization:
    organization = (
        PartnerOrganization.objects.select_for_update().filter(pk=run.organization_id).first()
    )
    if organization is None or organization.status != PartnerOrganization.Status.ACTIVE:
        raise ValidationError("legacy migration organization is not active")
    return organization


def _verified_dry_run_manifest(run: LegacyMigrationRun) -> LegacyMigrationManifest:
    manifest = (
        run.manifests.filter(kind=LegacyMigrationManifest.Kind.DRY_RUN).order_by("sequence").first()
    )
    if manifest is None:
        raise ValidationError("legacy migration has no dry-run manifest")
    if manifest.digest != _sha256(_canonical_bytes(manifest.payload)):
        raise ValidationError("dry-run manifest digest is invalid")
    if manifest.digest != run.dry_run_manifest_digest:
        raise ValidationError("dry-run manifest no longer matches the migration run")
    if (
        manifest.payload.get("run_id") != str(run.pk)
        or manifest.payload.get("plan_digest") != run.plan_digest
        or manifest.payload.get("source_state_digest") != run.source_state_digest
    ):
        raise ValidationError("dry-run manifest provenance no longer matches the migration run")
    persisted_payload = _persisted_preview_manifest_payload(run)
    if _sha256(_canonical_bytes(persisted_payload)) != manifest.digest:
        raise ValidationError("persisted preview evidence no longer matches the dry-run manifest")
    return manifest


def _verified_approval(
    run: LegacyMigrationRun, expected_manifest_digest: str
) -> LegacyMigrationDecisionRevision:
    manifest = _verified_dry_run_manifest(run)
    if expected_manifest_digest != manifest.digest:
        raise ValidationError("apply does not match the approved manifest digest")
    decision = run.decision_revisions.order_by("-sequence").first()
    if (
        decision is None
        or decision.decision != LegacyMigrationDecisionRevision.Decision.APPROVE
        or decision.manifest_digest != manifest.digest
    ):
        raise ValidationError("legacy migration is not approved for this exact manifest")
    return decision


@transaction.atomic
def approve_migration(
    *,
    actor: Account,
    run: LegacyMigrationRun,
    manifest_digest: str,
    rationale: str,
) -> LegacyMigrationRun:
    """Append an exact-digest, tenant-scoped reviewer approval."""

    _lock_organization(run)
    current = _current_run(run)
    _authorize_reviewer(actor, current)
    if current.state != LegacyMigrationRun.State.DRY_RUN_COMPLETE:
        raise ValidationError("only a completed dry run can be approved")
    manifest = _verified_dry_run_manifest(current)
    if manifest_digest != manifest.digest:
        raise ValidationError("approval must reference the exact dry-run manifest digest")
    if current.decision_revisions.exists():
        raise ValidationError("legacy migration already contains a reviewer decision")
    LegacyMigrationDecisionRevision.objects.create(
        organization=current.organization,
        run=current,
        sequence=1,
        decision=LegacyMigrationDecisionRevision.Decision.APPROVE,
        manifest_digest=manifest.digest,
        predecessor=None,
        actor=actor,
        rationale=_rationale(rationale),
    )
    current.state = LegacyMigrationRun.State.APPROVED
    current.save(update_fields=["state", "updated_at"])
    return current


def _chunk_payload(
    *,
    run: LegacyMigrationRun,
    sequence: int,
    table_plan: LegacyTablePlan,
    previews: Sequence[LegacyMigrationRowPreview],
) -> dict[str, object]:
    return {
        "version": "legacy-apply-checkpoint-v1",
        "run_id": str(run.pk),
        "dry_run_manifest_digest": run.dry_run_manifest_digest,
        "plan_digest": run.plan_digest,
        "source_state_digest": run.source_state_digest,
        "sequence": sequence,
        "sheet_index": table_plan.worksheet.sheet_index,
        "table_plan_id": str(table_plan.pk),
        "table_mapping_digest": table_plan.mapping_digest,
        "rows": [
            {
                "preview_id": str(preview.pk),
                "source_row_number": preview.source_row_number,
                "payload_digest": preview.payload_digest,
                "dry_run_classification": preview.classification,
            }
            for preview in previews
        ],
    }


def _planned_chunks(run: LegacyMigrationRun, chunk_size: int) -> list[_ApplyChunk]:
    previews = list(
        run.row_previews.select_related("table_plan__worksheet")
        .exclude(classification=LegacyMigrationRowPreview.Classification.IGNORED)
        .order_by(
            "source_sheet_index",
            "table_plan__start_row",
            "table_plan__start_column",
            "table_plan__label",
            "table_plan_id",
            "source_row_number",
            "preview_id",
        )
    )
    chunks: list[_ApplyChunk] = []
    sequence = 1
    offset = 0
    while offset < len(previews):
        table_plan_id = previews[offset].table_plan_id
        end = offset
        while end < len(previews) and previews[end].table_plan_id == table_plan_id:
            end += 1
        table_previews = previews[offset:end]
        for chunk_start in range(0, len(table_previews), chunk_size):
            selected = tuple(table_previews[chunk_start : chunk_start + chunk_size])
            table_plan = selected[0].table_plan
            payload = _chunk_payload(
                run=run,
                sequence=sequence,
                table_plan=table_plan,
                previews=selected,
            )
            chunks.append(
                _ApplyChunk(
                    sequence=sequence,
                    table_plan=table_plan,
                    previews=selected,
                    digest=_sha256(_canonical_bytes(payload)),
                )
            )
            sequence += 1
        offset = end
    return chunks


def _canonical_mapping_for_run(*, actor: Account, run: LegacyMigrationRun) -> MappingTemplate:
    mapping, _ = MappingTemplate.objects.get_or_create(
        organization=run.organization,
        name=f"Legacy canonical {run.pk}",
        version=1,
        defaults={
            "field_map": {field: field for field in _CANONICAL_RESULT_FIELDS},
            "created_by": actor,
        },
    )
    return mapping


def _assert_active_apply_claim(
    job: LegacyMigrationJob,
    *,
    claim_token: str,
    claim_generation: int,
    claim_owner: str,
    now: datetime | None = None,
) -> None:
    checked_at = now if now is not None else timezone.now()
    if (
        job.status != LegacyMigrationJob.Status.RUNNING
        or job.claim_token != claim_token
        or job.claim_generation != claim_generation
        or job.claim_owner != claim_owner
    ):
        raise ApplyClaimLost("legacy migration APPLY claim is stale or superseded")
    if job.lease_expires_at is None or job.lease_expires_at <= checked_at:
        raise ApplyClaimLost("legacy migration APPLY claim has expired")


@transaction.atomic
def _prepare_claimed_apply(
    *,
    actor: Account,
    run: LegacyMigrationRun,
    job: LegacyMigrationJob,
    approved_manifest_digest: str,
    claim_token: str,
    claim_generation: int,
    claim_owner: str,
) -> tuple[LegacyMigrationRun, MappingTemplate, list[_ApplyChunk], int]:
    _lock_organization(run)
    current = _current_run(run)
    _authorize(actor, current)
    if current.state == LegacyMigrationRun.State.WITHDRAWN:
        raise ValidationError("withdrawn legacy migrations cannot be applied")
    _verified_approval(current, approved_manifest_digest)
    if current.state != LegacyMigrationRun.State.APPLYING:
        raise ValidationError("legacy migration is not in applying state")
    current_job = LegacyMigrationJob.objects.select_for_update().get(pk=job.pk)
    if (
        current_job.run_id != current.pk
        or current_job.organization_id != current.organization_id
        or current_job.job_kind != LegacyMigrationJob.Kind.APPLY
        or current_job.requested_manifest_digest != approved_manifest_digest
    ):
        raise ValidationError("apply job provenance does not match the migration run")
    _assert_active_apply_claim(
        current_job,
        claim_token=claim_token,
        claim_generation=claim_generation,
        claim_owner=claim_owner,
    )
    mapping = _canonical_mapping_for_run(actor=actor, run=current)
    chunks = _planned_chunks(current, current_job.chunk_size)
    checkpoints = list(
        current.checkpoints.select_for_update()
        # PostgreSQL cannot apply FOR UPDATE through the nullable
        # checkpoint->ingestion_run outer join. Lock only the checkpoint and
        # its required plan relationships; committed ingestion evidence is
        # loaded lazily below while the organization/checkpoint locks remain
        # held.
        .select_related("worksheet", "table_plan")
        .order_by("sequence")
    )
    allowed_checkpoint_job_ids = _allowed_checkpoint_job_ids(run=current, job=current_job)
    if len(checkpoints) > len(chunks):
        raise ValidationError("checkpoint evidence exceeds the deterministic apply plan")
    for expected_sequence, checkpoint in enumerate(checkpoints, start=1):
        if checkpoint.sequence != expected_sequence:
            raise ValidationError("checkpoint evidence is not a contiguous committed prefix")
        chunk = chunks[expected_sequence - 1]
        _assert_checkpoint_matches(
            checkpoint,
            chunk,
            run=current,
            allowed_job_ids=allowed_checkpoint_job_ids,
        )
        if checkpoint.status != LegacyMigrationCheckpoint.Status.COMMITTED:
            raise ValidationError("checkpoint evidence is not a contiguous committed prefix")
        _assert_checkpoint_ingestion(run=current, checkpoint=checkpoint, chunk=chunk)
    next_sequence = len(checkpoints) + 1
    if current_job.next_checkpoint_sequence != next_sequence:
        current_job.next_checkpoint_sequence = next_sequence
        current_job.save(update_fields=["next_checkpoint_sequence", "updated_at"])
    return current, mapping, chunks, next_sequence


def _assert_checkpoint_matches(
    checkpoint: LegacyMigrationCheckpoint,
    chunk: _ApplyChunk,
    *,
    run: LegacyMigrationRun,
    allowed_job_ids: Collection[UUID],
) -> None:
    if (
        checkpoint.organization_id != run.organization_id
        or checkpoint.run_id != run.pk
        or checkpoint.job_id not in allowed_job_ids
        or checkpoint.table_plan_id != chunk.table_plan.pk
        or checkpoint.worksheet_id != chunk.table_plan.worksheet_id
        or checkpoint.first_source_row != chunk.previews[0].source_row_number
        or checkpoint.last_source_row != chunk.previews[-1].source_row_number
        or checkpoint.checkpoint_digest != chunk.digest
    ):
        raise ValidationError("existing checkpoint does not match the deterministic apply plan")


def _allowed_checkpoint_job_ids(
    *,
    run: LegacyMigrationRun,
    job: LegacyMigrationJob,
) -> set[UUID]:
    predecessor_statuses = (
        LegacyMigrationJob.Status.FAILED,
        LegacyMigrationJob.Status.CANCELLED,
    )
    return set(
        LegacyMigrationJob.objects.filter(
            organization_id=run.organization_id,
            run_id=run.pk,
            job_kind=LegacyMigrationJob.Kind.APPLY,
            requested_manifest_digest=job.requested_manifest_digest,
            chunk_size=job.chunk_size,
            job_sequence__lte=job.job_sequence,
        )
        .filter(models.Q(pk=job.pk) | models.Q(status__in=predecessor_statuses))
        .values_list("job_id", flat=True)
    )


def _assert_checkpoint_ingestion(
    *,
    run: LegacyMigrationRun,
    checkpoint: LegacyMigrationCheckpoint,
    chunk: _ApplyChunk,
) -> None:
    ingestion = checkpoint.ingestion_run
    expected_partition = f"legacy:{run.pk}:{chunk.sequence}:{chunk.digest[:32]}"
    if (
        ingestion is None
        or ingestion.organization_id != run.organization_id
        or ingestion.artifact_id != run.inventory.artifact_id
        or ingestion.mapping_template.organization_id != run.organization_id
        or ingestion.source_sheet_name != chunk.table_plan.worksheet.sheet_name
        or ingestion.source_partition_key != expected_partition
        or ingestion.total_rows != len(chunk.previews)
        or ingestion.total_rows
        != (
            ingestion.published_rows
            + ingestion.quarantined_rows
            + ingestion.duplicate_rows
            + ingestion.conflict_rows
        )
    ):
        raise ValidationError("committed checkpoint ingestion evidence is invalid")
    staged_results = list(ingestion.staged_results.order_by("row_number"))
    if len(staged_results) != len(chunk.previews):
        raise ValidationError("committed checkpoint ingestion evidence is invalid")
    for staged, preview in zip(staged_results, chunk.previews, strict=True):
        if (
            staged.source_sheet_name != ingestion.source_sheet_name
            or staged.source_row_number != preview.source_row_number
            or staged.payload_digest != preview.payload_digest
        ):
            raise ValidationError("committed checkpoint ingestion evidence is invalid")


@transaction.atomic
def _commit_apply_chunk(
    *,
    actor: Account,
    run: LegacyMigrationRun,
    job: LegacyMigrationJob,
    mapping: MappingTemplate,
    chunk: _ApplyChunk,
    approved_manifest_digest: str,
    claim_token: str,
    claim_generation: int,
    claim_owner: str,
) -> tuple[LegacyMigrationCheckpoint, bool]:
    """Commit one ordinary ingestion and its checkpoint as one outer transaction."""

    _lock_organization(run)
    current = _current_run(run)
    _authorize(actor, current)
    if current.state == LegacyMigrationRun.State.WITHDRAWN:
        raise ValidationError("withdrawn legacy migrations cannot be applied")
    if current.state != LegacyMigrationRun.State.APPLYING:
        raise ValidationError("legacy migration is not in applying state")
    _verified_approval(current, approved_manifest_digest)
    current_job = LegacyMigrationJob.objects.select_for_update().get(pk=job.pk)
    if current_job.run_id != current.pk or current_job.organization_id != current.organization_id:
        raise ValidationError("apply job does not belong to the migration run")
    _assert_active_apply_claim(
        current_job,
        claim_token=claim_token,
        claim_generation=claim_generation,
        claim_owner=claim_owner,
    )
    existing = current.checkpoints.select_for_update().filter(sequence=chunk.sequence).first()
    if existing is not None:
        allowed_checkpoint_job_ids = _allowed_checkpoint_job_ids(
            run=current,
            job=current_job,
        )
        _assert_checkpoint_matches(
            existing,
            chunk,
            run=current,
            allowed_job_ids=allowed_checkpoint_job_ids,
        )
        if existing.status == LegacyMigrationCheckpoint.Status.COMMITTED:
            _assert_checkpoint_ingestion(run=current, checkpoint=existing, chunk=chunk)
            if current_job.next_checkpoint_sequence < chunk.sequence + 1:
                current_job.next_checkpoint_sequence = chunk.sequence + 1
                current_job.save(update_fields=["next_checkpoint_sequence", "updated_at"])
            return existing, False

    partition_key = f"legacy:{current.pk}:{chunk.sequence}:{chunk.digest[:32]}"
    outcome = ingest_preparsed_rows(
        actor=actor,
        organization=current.organization,
        mapping_template=mapping,
        operator_key=partition_key,
        rows=[dict(preview.canonical_payload) for preview in chunk.previews],
        artifact_kind=SourceArtifact.Kind.SPREADSHEET,
        artifact_digest=current.inventory.artifact.digest,
        artifact_name=current.inventory.artifact.original_name,
        content_type=current.inventory.artifact.content_type,
        byte_size=current.inventory.artifact.byte_size,
        object_reference=current.inventory.artifact.object_reference,
        source_sheet_name=chunk.table_plan.worksheet.sheet_name,
        source_partition_key=partition_key,
        source_row_numbers=[preview.source_row_number for preview in chunk.previews],
    )
    now = timezone.now()
    values = {
        "status": LegacyMigrationCheckpoint.Status.COMMITTED,
        "ingestion_run": outcome.run,
        "total_rows": outcome.run.total_rows,
        "published_rows": outcome.run.published_rows,
        "quarantined_rows": outcome.run.quarantined_rows,
        "duplicate_rows": outcome.run.duplicate_rows,
        "conflict_rows": outcome.run.conflict_rows,
        "completed_at": now,
    }
    if existing is None:
        checkpoint = LegacyMigrationCheckpoint.objects.create(
            organization=current.organization,
            run=current,
            job=job,
            worksheet=chunk.table_plan.worksheet,
            table_plan=chunk.table_plan,
            sequence=chunk.sequence,
            first_source_row=chunk.previews[0].source_row_number,
            last_source_row=chunk.previews[-1].source_row_number,
            checkpoint_digest=chunk.digest,
            **values,
        )
    else:
        for field, value in values.items():
            setattr(existing, field, value)
        existing.save(update_fields=[*values])
        checkpoint = existing

    current_job.next_checkpoint_sequence = chunk.sequence + 1
    current_job.save(update_fields=["next_checkpoint_sequence", "updated_at"])
    return checkpoint, True


def _reconciliation_payload(
    *,
    run: LegacyMigrationRun,
    job: LegacyMigrationJob,
    chunks: Sequence[_ApplyChunk],
    checkpoints: Sequence[LegacyMigrationCheckpoint],
) -> dict[str, object]:
    allowed_checkpoint_job_ids = _allowed_checkpoint_job_ids(run=run, job=job)
    checkpoint_by_sequence = {checkpoint.sequence: checkpoint for checkpoint in checkpoints}
    row_evidence: dict[str, dict[str, object]] = {}
    totals = {
        "published_rows": 0,
        "quarantined_rows": 0,
        "duplicate_rows": 0,
        "conflict_rows": 0,
    }
    checkpoint_evidence: list[dict[str, object]] = []
    for chunk in chunks:
        checkpoint = checkpoint_by_sequence.get(chunk.sequence)
        if checkpoint is None or checkpoint.status != LegacyMigrationCheckpoint.Status.COMMITTED:
            raise ValidationError("reconciliation requires every apply checkpoint")
        _assert_checkpoint_matches(
            checkpoint,
            chunk,
            run=run,
            allowed_job_ids=allowed_checkpoint_job_ids,
        )
        ingestion_run = checkpoint.ingestion_run
        _assert_checkpoint_ingestion(run=run, checkpoint=checkpoint, chunk=chunk)
        assert ingestion_run is not None
        staged_results = list(ingestion_run.staged_results.order_by("row_number"))
        staged_by_source_row = {staged.source_row_number: staged for staged in staged_results}
        if len(staged_results) != len(chunk.previews) or len(staged_by_source_row) != len(
            chunk.previews
        ):
            raise ValidationError("checkpoint staged rows do not match preview coordinates")
        published_by_staged = {
            published.staged_result_id: published
            for published in PublishedSourceResult.objects.filter(
                organization=run.organization,
                artifact=run.inventory.artifact,
                staged_result_id__in=[staged.pk for staged in staged_results],
            )
        }
        for preview in chunk.previews:
            staged = staged_by_source_row.get(preview.source_row_number)
            if staged is None:
                raise ValidationError("checkpoint omitted a preview source row")
            published = published_by_staged.get(staged.pk)
            if staged.outcome == StagedResult.Outcome.PUBLISHED and published is None:
                raise ValidationError("published checkpoint row has invalid tenant provenance")
            row_evidence[str(preview.pk)] = {
                "preview_id": str(preview.pk),
                "sheet_index": preview.source_sheet_index,
                "table_plan_id": str(preview.table_plan_id),
                "source_row_number": preview.source_row_number,
                "payload_digest": preview.payload_digest,
                "dry_run_classification": preview.classification,
                "apply_outcome": staged.outcome,
                "diverged": preview.classification != staged.outcome,
                "checkpoint_sequence": checkpoint.sequence,
                "ingestion_run_id": str(ingestion_run.pk),
                "staged_result_id": str(staged.pk),
                "published_result_id": (str(published.pk) if published is not None else ""),
            }
        totals["published_rows"] += checkpoint.published_rows
        totals["quarantined_rows"] += checkpoint.quarantined_rows
        totals["duplicate_rows"] += checkpoint.duplicate_rows
        totals["conflict_rows"] += checkpoint.conflict_rows
        checkpoint_evidence.append(
            {
                "sequence": checkpoint.sequence,
                "sheet_index": checkpoint.worksheet.sheet_index,
                "table_plan_id": str(checkpoint.table_plan_id),
                "first_source_row": checkpoint.first_source_row,
                "last_source_row": checkpoint.last_source_row,
                "checkpoint_digest": checkpoint.checkpoint_digest,
                "ingestion_run_id": str(ingestion_run.pk),
                "total_rows": checkpoint.total_rows,
                "published_rows": checkpoint.published_rows,
                "quarantined_rows": checkpoint.quarantined_rows,
                "duplicate_rows": checkpoint.duplicate_rows,
                "conflict_rows": checkpoint.conflict_rows,
            }
        )

    previews = list(
        run.row_previews.select_related("table_plan").order_by(
            "source_sheet_index",
            "source_row_number",
            "source_column_start",
            "preview_id",
        )
    )
    for preview in previews:
        key = str(preview.pk)
        if preview.classification == LegacyMigrationRowPreview.Classification.IGNORED:
            row_evidence[key] = {
                "preview_id": key,
                "sheet_index": preview.source_sheet_index,
                "table_plan_id": str(preview.table_plan_id),
                "source_row_number": preview.source_row_number,
                "payload_digest": preview.payload_digest,
                "dry_run_classification": preview.classification,
                "apply_outcome": "ignored",
                "diverged": False,
                "checkpoint_sequence": None,
                "ingestion_run_id": "",
                "staged_result_id": "",
                "published_result_id": "",
            }
        elif key not in row_evidence:
            raise ValidationError("reconciliation omitted a candidate preview row")

    worksheets = list(run.worksheets.prefetch_related("table_plans").order_by("sheet_index"))
    sheet_evidence = [
        {
            "sheet_index": worksheet.sheet_index,
            "visibility": worksheet.visibility,
            "disposition": worksheet.disposition,
            "metadata_digest": worksheet.metadata_digest,
            "table_count": len(worksheet.table_plans.all()),
            "preview_rows": sum(
                preview.source_sheet_index == worksheet.sheet_index for preview in previews
            ),
        }
        for worksheet in worksheets
    ]
    ignored_rows = sum(
        preview.classification == LegacyMigrationRowPreview.Classification.IGNORED
        for preview in previews
    )
    complete_totals: dict[str, int] = {
        "discovered_sheets": len(worksheets),
        "included_sheets": sum(
            worksheet.disposition == LegacyWorksheetPlan.Disposition.INCLUDED
            for worksheet in worksheets
        ),
        "ignored_sheets": sum(
            worksheet.disposition == LegacyWorksheetPlan.Disposition.IGNORED
            for worksheet in worksheets
        ),
        "preview_rows": len(previews),
        "candidate_rows": len(previews) - ignored_rows,
        "ignored_rows": ignored_rows,
        **totals,
    }
    ordered_rows = [row_evidence[str(preview.pk)] for preview in previews]
    return {
        "version": "legacy-reconciliation-manifest-v1",
        "run_id": str(run.pk),
        "dry_run_manifest_digest": run.dry_run_manifest_digest,
        "plan_digest": run.plan_digest,
        "source_state_digest": run.source_state_digest,
        "totals": complete_totals,
        "sheets": sheet_evidence,
        "checkpoints": checkpoint_evidence,
        "rows": ordered_rows,
    }


@transaction.atomic
def _complete_apply(
    *,
    actor: Account,
    run: LegacyMigrationRun,
    job: LegacyMigrationJob,
    chunks: Sequence[_ApplyChunk],
    approved_manifest_digest: str,
    claim_token: str,
    claim_generation: int,
    claim_owner: str,
) -> LegacyMigrationRun:
    _lock_organization(run)
    current = _current_run(run)
    _authorize(actor, current)
    if current.state == LegacyMigrationRun.State.WITHDRAWN:
        raise ValidationError("withdrawn legacy migrations cannot be completed")
    _verified_approval(current, approved_manifest_digest)
    current_job = LegacyMigrationJob.objects.select_for_update().get(pk=job.pk)
    if current_job.run_id != current.pk or current_job.organization_id != current.organization_id:
        raise ValidationError("apply job does not belong to the migration run")
    _assert_active_apply_claim(
        current_job,
        claim_token=claim_token,
        claim_generation=claim_generation,
        claim_owner=claim_owner,
    )
    checkpoints = list(
        current.checkpoints.select_related("worksheet", "table_plan", "ingestion_run").order_by(
            "sequence"
        )
    )
    if len(checkpoints) != len(chunks):
        raise ValidationError("apply cannot complete until every deterministic chunk is committed")
    payload = _reconciliation_payload(
        run=current,
        job=current_job,
        chunks=chunks,
        checkpoints=checkpoints,
    )
    digest = _sha256(_canonical_bytes(payload))
    existing_manifest = current.manifests.filter(
        kind=LegacyMigrationManifest.Kind.RECONCILIATION
    ).first()
    if existing_manifest is not None:
        if existing_manifest.digest != digest:
            raise ValidationError("existing reconciliation manifest does not match apply evidence")
    else:
        LegacyMigrationManifest.objects.create(
            organization=current.organization,
            run=current,
            kind=LegacyMigrationManifest.Kind.RECONCILIATION,
            sequence=1,
            digest=digest,
            payload=payload,
            created_by=actor,
        )
    current_job.status = LegacyMigrationJob.Status.COMPLETED
    current_job.next_checkpoint_sequence = len(chunks) + 1
    current_job.completed_at = timezone.now()
    current_job.last_error_code = ""
    current_job.last_error_message = ""
    current_job.claim_token = ""
    current_job.claim_owner = ""
    current_job.heartbeat_at = None
    current_job.lease_expires_at = None
    current_job.save(
        update_fields=[
            "status",
            "next_checkpoint_sequence",
            "completed_at",
            "last_error_code",
            "last_error_message",
            "claim_token",
            "claim_owner",
            "heartbeat_at",
            "lease_expires_at",
            "updated_at",
        ]
    )
    current.applied_rows = int(payload["totals"]["candidate_rows"])  # type: ignore[index]
    current.state = LegacyMigrationRun.State.COMPLETED
    current.save(update_fields=["applied_rows", "state", "updated_at"])
    return current


def apply_migration(
    *,
    actor: Account,
    run: LegacyMigrationRun,
    workbook_content: bytes,
    approved_manifest_digest: str,
    chunk_size: int = MAX_APPLY_CHUNK_ROWS,
    request_rationale: str = "Authorized synchronous compatibility APPLY request.",
    fail_after_committed_chunks: int | None = None,
) -> LegacyMigrationRun:
    """Synchronously orchestrate the durable APPLY queue for compatibility callers."""

    if (
        isinstance(chunk_size, bool)
        or not isinstance(chunk_size, int)
        or not (1 <= chunk_size <= MAX_APPLY_CHUNK_ROWS)
    ):
        raise ValidationError(f"chunk size must be between 1 and {MAX_APPLY_CHUNK_ROWS}")
    if fail_after_committed_chunks is not None:
        fail_after_committed_chunks = _positive_integer(
            fail_after_committed_chunks,
            field="failure injection count",
        )

    # Deny before parsing or comparing caller-supplied source bytes so an
    # unauthorized account cannot use validation errors as a source oracle.
    _authorize(actor, run)
    inspection = _verify_source(run, workbook_content)
    _verify_persisted_plan(run, inspection)

    from mnemex.legacy_migration.jobs import (
        _claim_specific_apply_job,
        enqueue_apply_job,
        execute_apply_job,
    )

    job = enqueue_apply_job(
        actor=actor,
        run=run,
        approved_manifest_digest=approved_manifest_digest,
        chunk_size=chunk_size,
        request_rationale=request_rationale,
    )
    if job.status == LegacyMigrationJob.Status.COMPLETED:
        return LegacyMigrationRun.objects.get(pk=run.pk)
    claim = _claim_specific_apply_job(
        job_id=job.pk,
        owner=f"sync-{actor.pk}",
        allow_early_retry=True,
    )
    if claim is None:
        raise ValidationError("apply job is currently owned by another worker")
    execute_apply_job(
        claim,
        workbook_content=workbook_content,
        fail_after_committed_chunks=fail_after_committed_chunks,
        _raise_error=True,
    )
    return LegacyMigrationRun.objects.get(pk=run.pk)


@transaction.atomic
def withdraw_migration(
    *,
    actor: Account,
    run: LegacyMigrationRun,
    manifest_digest: str,
    rationale: str,
) -> LegacyMigrationRun:
    """Append a logical withdrawal without deleting source or publication evidence."""

    _lock_organization(run)
    current = _current_run(run)
    _authorize_reviewer(actor, current)
    if current.state == LegacyMigrationRun.State.WITHDRAWN:
        raise ValidationError("legacy migration is already withdrawn")
    if current.state not in {
        LegacyMigrationRun.State.APPROVED,
        LegacyMigrationRun.State.APPLYING,
        LegacyMigrationRun.State.FAILED,
        LegacyMigrationRun.State.COMPLETED,
    }:
        raise ValidationError("legacy migration must be approved before withdrawal")
    reconciliation = (
        current.manifests.filter(kind=LegacyMigrationManifest.Kind.RECONCILIATION)
        .order_by("-sequence")
        .first()
    )
    authority_digest = (
        reconciliation.digest if reconciliation is not None else current.dry_run_manifest_digest
    )
    if manifest_digest != authority_digest:
        raise ValidationError("withdrawal must reference the exact current manifest digest")
    predecessor = current.decision_revisions.order_by("-sequence").first()
    if (
        predecessor is None
        or predecessor.decision != LegacyMigrationDecisionRevision.Decision.APPROVE
    ):
        raise ValidationError("withdrawal requires an active reviewer approval")
    decision = LegacyMigrationDecisionRevision.objects.create(
        organization=current.organization,
        run=current,
        sequence=predecessor.sequence + 1,
        decision=LegacyMigrationDecisionRevision.Decision.WITHDRAW,
        manifest_digest=authority_digest,
        predecessor=predecessor,
        actor=actor,
        rationale=_rationale(rationale),
    )

    partition_prefix = f"legacy:{current.pk}:"
    ingestion_runs = list(
        IngestionRun.objects.filter(
            organization=current.organization,
            source_partition_key__startswith=partition_prefix,
        ).order_by("created_at", "run_id")
    )
    staged_ids = list(
        StagedResult.objects.filter(run__in=ingestion_runs).values_list(
            "staged_result_id", flat=True
        )
    )
    published_ids = list(
        PublishedSourceResult.objects.filter(staged_result_id__in=staged_ids)
        .order_by("published_at", "published_result_id")
        .values_list("published_result_id", flat=True)
    )
    checkpoints = list(current.checkpoints.order_by("sequence"))
    payload = {
        "version": "legacy-withdrawal-manifest-v1",
        "run_id": str(current.pk),
        "authority_manifest_digest": authority_digest,
        "decision_id": str(decision.pk),
        "predecessor_decision_id": str(predecessor.pk),
        "source_artifact_id": str(current.inventory.artifact_id),
        "checkpoint_ids": [str(checkpoint.pk) for checkpoint in checkpoints],
        "ingestion_run_ids": [str(ingestion_run.pk) for ingestion_run in ingestion_runs],
        "staged_result_ids": [str(staged_id) for staged_id in staged_ids],
        "published_result_ids": [str(published_id) for published_id in published_ids],
    }
    LegacyMigrationManifest.objects.create(
        organization=current.organization,
        run=current,
        kind=LegacyMigrationManifest.Kind.WITHDRAWAL,
        sequence=1,
        digest=_sha256(_canonical_bytes(payload)),
        payload=payload,
        created_by=actor,
    )
    active_jobs = list(
        current.jobs.select_for_update().filter(
            job_kind=LegacyMigrationJob.Kind.APPLY,
            status__in={
                LegacyMigrationJob.Status.PENDING,
                LegacyMigrationJob.Status.RUNNING,
                LegacyMigrationJob.Status.RETRY_WAIT,
            },
        )
    )
    cancelled_at = timezone.now()
    for job in active_jobs:
        job.status = LegacyMigrationJob.Status.CANCELLED
        job.claim_token = ""
        job.claim_owner = ""
        job.heartbeat_at = None
        job.lease_expires_at = None
        job.completed_at = cancelled_at
        job.last_error_code = "authorization_withdrawn"
        job.last_error_message = "migration authorization was withdrawn"
        job.save(
            update_fields=[
                "status",
                "claim_token",
                "claim_owner",
                "heartbeat_at",
                "lease_expires_at",
                "completed_at",
                "last_error_code",
                "last_error_message",
                "updated_at",
            ]
        )
    current.state = LegacyMigrationRun.State.WITHDRAWN
    current.save(update_fields=["state", "updated_at"])
    return current
