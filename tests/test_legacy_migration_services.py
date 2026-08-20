from __future__ import annotations

import hashlib
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from django.core.exceptions import ValidationError
from openpyxl import Workbook

from mnemex.accounts.models import Account, PrivilegedRoleAssignment
from mnemex.legacy_migration.inspection import inspect_legacy_workbook
from mnemex.legacy_migration.models import (
    LegacyMigrationManifest,
    LegacyMigrationRowPreview,
    LegacyMigrationRun,
    LegacySourceInventory,
    LegacyTablePlan,
    LegacyWorksheetPlan,
)
from mnemex.legacy_migration.services import (
    configure_migration_run,
    dry_run_migration,
    stable_artifact_coordinate_id,
)
from mnemex.partners.models import PartnerOrganization
from mnemex.results.models import (
    IngestionRun,
    MappingTemplate,
    PublishedSourceResult,
    ReconciliationCase,
    SourceArtifact,
    StagedResult,
)
from mnemex.results.services import ingest_manual_rows
from tests.mfa_helpers import enroll_account_mfa

pytestmark = pytest.mark.django_db


CANONICAL_HEADERS = ["Result ID", "Competitor", "Score"]


def _actor(organization: PartnerOrganization) -> Account:
    actor = Account.objects.create_user(
        email=f"legacy-{organization.pk}@mnemex.example.invalid",
        password="synthetic-password-123",
    )
    enroll_account_mfa(actor)
    PrivilegedRoleAssignment.objects.create(
        account=actor,
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
        scope=PrivilegedRoleAssignment.Scope.ORGANIZATION,
        organization=organization,
        assigned_by=actor,
    )
    return actor


def _xlsx(
    sheets: list[tuple[str, str]],
    *,
    rows_by_sheet: dict[str, list[list[object]]] | None = None,
    headers_by_sheet: dict[str, list[object]] | None = None,
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, (name, visibility) in enumerate(sheets):
        sheet = workbook.create_sheet(title=name)
        sheet.sheet_state = visibility
        sheet.merge_cells("A1:C1")
        sheet["A1"] = f"Synthetic table {index + 1}"
        headers = (headers_by_sheet or {}).get(name, CANONICAL_HEADERS)
        for column, value in enumerate(headers, start=1):
            sheet.cell(row=5, column=column, value=value)
        for row_number, values in enumerate((rows_by_sheet or {}).get(name, []), start=6):
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row_number, column=column, value=value)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _with_extra_archive_entries(content: bytes, *, count: int) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry, source.read(entry.filename))
        for index in range(count):
            target.writestr(f"synthetic-extra/entry-{index:03}.txt", "synthetic")
    return output.getvalue()


def _inventory_run(
    content: bytes,
    organization: PartnerOrganization,
    actor: Account,
    *,
    run_version: int = 1,
) -> LegacyMigrationRun:
    digest = hashlib.sha256(content).hexdigest()
    artifact, _ = SourceArtifact.objects.get_or_create(
        organization=organization,
        kind=SourceArtifact.Kind.SPREADSHEET,
        digest=digest,
        defaults={
            "original_name": "synthetic-legacy.xlsx",
            "content_type": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            "byte_size": len(content),
            "object_reference": "synthetic/private/legacy.xlsx",
            "uploaded_by": actor,
        },
    )
    inventory = LegacySourceInventory.objects.create(
        organization=organization,
        artifact=artifact,
        source_key=f"synthetic-legacy-source-{run_version}",
        data_rights_reference="synthetic-test-authorization",
        parser_version="legacy-xlsx-inspector-v1",
        manifest_digest=hashlib.sha256(f"inventory-{run_version}".encode()).hexdigest(),
        created_by=actor,
    )
    return LegacyMigrationRun.objects.create(
        organization=organization,
        inventory=inventory,
        run_version=run_version,
        ruleset_version="legacy-mapping-v1",
        created_by=actor,
    )


def _rules(*, derived_id: bool = False) -> dict[str, dict[str, object]]:
    source_id: dict[str, object]
    if derived_id:
        source_id = {"kind": "derived", "name": "artifact_coordinate_id"}
    else:
        source_id = {"kind": "column", "column": "Result ID"}
    return {
        "source_result_id": source_id,
        "source_revision": {"kind": "constant", "value": 1},
        "source_event_id": {"kind": "constant", "value": "synthetic-event-2026"},
        "event_name": {"kind": "constant", "value": "Synthetic Legacy Show"},
        "result_date": {"kind": "constant", "value": "2026-07-18"},
        "competitor_name": {"kind": "column", "column": "Competitor"},
        "discipline": {"kind": "constant", "value": "UNDERHAND"},
        "score_type": {"kind": "constant", "value": "time"},
        "score": {"kind": "column", "column": "Score"},
    }


def _configuration(
    *,
    included_name: str,
    ignored_name: str | None = None,
    rules: dict[str, dict[str, object]] | None = None,
    end_row: int = 10,
) -> list[dict[str, object]]:
    configuration: list[dict[str, object]] = [
        {
            "sheet_name": included_name,
            "disposition": "included",
            "tables": [
                {
                    "label": "qualifying",
                    "header_row": 5,
                    "start_row": 5,
                    "end_row": end_row,
                    "start_column": 1,
                    "end_column": 3,
                    "mapping_rules": rules or _rules(),
                }
            ],
        }
    ]
    if ignored_name is not None:
        configuration.append(
            {
                "sheet_name": ignored_name,
                "disposition": "ignored",
                "ignore_reason": "Synthetic summary is not row-level results.",
                "tables": [],
            }
        )
    return configuration


def _identity_mapping(organization: PartnerOrganization, actor: Account) -> MappingTemplate:
    fields = (
        "source_result_id",
        "source_revision",
        "source_event_id",
        "event_name",
        "result_date",
        "competitor_name",
        "discipline",
        "score_type",
        "score",
    )
    return MappingTemplate.objects.create(
        organization=organization,
        name="Synthetic seed identity",
        version=1,
        field_map={field: field for field in fields},
        created_by=actor,
    )


def _canonical_row(result_id: str, competitor: str, score: str) -> dict[str, object]:
    return {
        "source_result_id": result_id,
        "source_revision": 1,
        "source_event_id": "synthetic-event-2026",
        "event_name": "Synthetic Legacy Show",
        "result_date": "2026-07-18",
        "competitor_name": competitor,
        "discipline": "UNDERHAND",
        "score_type": "time",
        "score": score,
    }


def test_discovery_supports_legacy_scale_and_preserves_sheet_metadata() -> None:
    sheets = [(f"Synthetic Sheet {index:02}", "visible") for index in range(42)]
    sheets[1] = (sheets[1][0], "hidden")
    sheets[2] = (sheets[2][0], "veryHidden")
    sheets[-1] = ("Synthetic Sheet 41 ", "visible")
    content = _with_extra_archive_entries(_xlsx(sheets), count=110)

    discovered = inspect_legacy_workbook(content)
    replay = inspect_legacy_workbook(content)

    assert len(discovered.sheets) == 42
    assert [sheet.name for sheet in discovered.sheets] == [name for name, _ in sheets]
    assert [sheet.visibility for sheet in discovered.sheets[:3]] == [
        "visible",
        "hidden",
        "very_hidden",
    ]
    assert 5 in discovered.sheets[0].header_candidates
    assert discovered.source_state_digest == replay.source_state_digest
    assert [sheet.metadata_digest for sheet in discovered.sheets] == [
        sheet.metadata_digest for sheet in replay.sheets
    ]


@pytest.mark.parametrize("sheet_count", [28, 38, 42])
def test_discovery_accepts_calibration_sheet_families(sheet_count: int) -> None:
    content = _xlsx([(f"Synthetic {index:02}", "visible") for index in range(sheet_count)])

    assert len(inspect_legacy_workbook(content).sheets) == sheet_count


def test_discovery_fails_closed_for_corrupt_and_excessive_workbooks() -> None:
    with pytest.raises(ValidationError, match="valid workbook archive"):
        inspect_legacy_workbook(b"not-an-xlsx")

    too_many_sheets = _xlsx([(f"Synthetic Excess {index:02}", "visible") for index in range(65)])
    with pytest.raises(ValidationError, match="sheet limit"):
        inspect_legacy_workbook(too_many_sheets)

    workbook = Workbook()
    workbook.active.cell(row=1_000, column=256, value="synthetic sparse corner")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    sparse_far_corner = stream.getvalue()
    assert len(sparse_far_corner) < 20_000
    with pytest.raises(ValidationError, match="cell-work limit"):
        inspect_legacy_workbook(sparse_far_corner)


def test_configuration_persists_exact_dispositions_tables_and_late_headers() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Legacy Configuration")
    actor = _actor(organization)
    content = _xlsx(
        [("Synthetic Qualifiers", "visible"), ("Synthetic Summary ", "hidden")],
        rows_by_sheet={"Synthetic Qualifiers": [["new-1", "Synthetic One", "12.34"]]},
        headers_by_sheet={"Synthetic Summary ": []},
    )
    run = _inventory_run(content, organization, actor)
    config = _configuration(
        included_name="Synthetic Qualifiers", ignored_name="Synthetic Summary ", end_row=6
    )

    configured = configure_migration_run(
        actor=actor,
        run=run,
        workbook_content=content,
        sheet_configurations=config,
    )

    assert configured.state == LegacyMigrationRun.State.CONFIGURED
    assert list(
        configured.worksheets.values_list("sheet_index", "sheet_name", "visibility", "disposition")
    ) == [
        (0, "Synthetic Qualifiers", "visible", "included"),
        (1, "Synthetic Summary ", "hidden", "ignored"),
    ]
    assert set(configured.table_plans.values_list("label", "header_row")) == {("qualifying", 5)}
    assert configured.worksheets.get(sheet_name="Synthetic Summary ").header_candidates == []
    assert configured.plan_digest
    assert configured.source_state_digest


def test_unknown_rules_and_duplicate_or_blank_headers_fail_without_partial_configuration() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Invalid Configuration")
    actor = _actor(organization)
    content = _xlsx(
        [("Synthetic Invalid", "visible")],
        headers_by_sheet={"Synthetic Invalid": ["Result ID", "", "Result ID"]},
    )
    run = _inventory_run(content, organization, actor)
    bad_rules = _rules()
    bad_rules["score"] = {"kind": "lookup", "column": "Score"}

    with pytest.raises(ValidationError, match="unknown mapping rule kind"):
        configure_migration_run(
            actor=actor,
            run=run,
            workbook_content=content,
            sheet_configurations=_configuration(included_name="Synthetic Invalid", rules=bad_rules),
        )

    assert LegacyWorksheetPlan.objects.filter(run=run).count() == 0
    assert LegacyTablePlan.objects.filter(run=run).count() == 0

    duplicate_content = _xlsx(
        [("Synthetic Duplicates", "visible")],
        headers_by_sheet={"Synthetic Duplicates": ["Result ID", "Competitor", "Competitor"]},
    )
    duplicate_run = _inventory_run(
        duplicate_content,
        organization,
        actor,
        run_version=2,
    )
    with pytest.raises(ValidationError, match="header names must be unique"):
        configure_migration_run(
            actor=actor,
            run=duplicate_run,
            workbook_content=duplicate_content,
            sheet_configurations=_configuration(included_name="Synthetic Duplicates"),
        )

    assert LegacyWorksheetPlan.objects.filter(run=duplicate_run).count() == 0
    assert LegacyTablePlan.objects.filter(run=duplicate_run).count() == 0
    run.refresh_from_db()
    assert run.state == LegacyMigrationRun.State.DISCOVERED

    with pytest.raises(ValidationError, match="header names cannot be blank"):
        configure_migration_run(
            actor=actor,
            run=run,
            workbook_content=content,
            sheet_configurations=_configuration(included_name="Synthetic Invalid"),
        )

    assert LegacyWorksheetPlan.objects.filter(run=run).count() == 0
    assert LegacyTablePlan.objects.filter(run=run).count() == 0


def test_configuration_rejects_overlapping_and_excessive_table_regions() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Region Bounds")
    actor = _actor(organization)
    content = _xlsx([("Synthetic Overlap", "visible")])
    run = _inventory_run(content, organization, actor)
    overlap = _configuration(included_name="Synthetic Overlap")
    overlap[0]["tables"].append(  # type: ignore[union-attr]
        {
            "label": "overlapping-final",
            "header_row": 5,
            "start_row": 5,
            "end_row": 10,
            "start_column": 1,
            "end_column": 3,
            "mapping_rules": _rules(derived_id=True),
        }
    )

    with pytest.raises(ValidationError, match="overlap"):
        configure_migration_run(
            actor=actor,
            run=run,
            workbook_content=content,
            sheet_configurations=overlap,
        )

    assert not run.table_plans.exists()

    excessive_content = _xlsx([("Synthetic Excess Regions", "visible")])
    excessive_run = _inventory_run(
        excessive_content,
        organization,
        actor,
        run_version=2,
    )
    excessive = _configuration(included_name="Synthetic Excess Regions")
    base_table = excessive[0]["tables"][0]  # type: ignore[index]
    excessive[0]["tables"] = [  # type: ignore[index]
        {**base_table, "label": f"region-{index:02}"} for index in range(33)
    ]

    with pytest.raises(ValidationError, match="table-region limit"):
        configure_migration_run(
            actor=actor,
            run=excessive_run,
            workbook_content=excessive_content,
            sheet_configurations=excessive,
        )

    assert not excessive_run.table_plans.exists()


def test_stable_coordinate_ids_are_retry_stable_and_coordinate_scoped() -> None:
    common = {
        "artifact_digest": "a" * 64,
        "sheet_index": 3,
        "sheet_name": "Synthetic Sheet ",
        "table_digest": "b" * 64,
        "source_row_number": 17,
    }

    first = stable_artifact_coordinate_id(**common)

    assert first == stable_artifact_coordinate_id(**common)
    assert first != stable_artifact_coordinate_id(**{**common, "artifact_digest": "c" * 64})
    assert first != stable_artifact_coordinate_id(**{**common, "sheet_index": 4})
    assert first != stable_artifact_coordinate_id(**{**common, "source_row_number": 18})
    assert len(first) <= 200


def test_side_by_side_tables_derive_distinct_retry_stable_source_ids() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Side By Side")
    actor = _actor(organization)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic Side By Side"
    for start_column in (1, 5):
        for offset, header in enumerate(CANONICAL_HEADERS):
            sheet.cell(row=5, column=start_column + offset, value=header)
    for start_column, result_id, competitor, score in (
        (1, "ignored-source-a", "Synthetic Qualifier", "12.3"),
        (5, "ignored-source-b", "Synthetic Final", "11.8"),
    ):
        for offset, value in enumerate((result_id, competitor, score)):
            sheet.cell(row=6, column=start_column + offset, value=value)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    content = stream.getvalue()
    run = _inventory_run(content, organization, actor)
    tables = []
    for label, start_column in (("qualifying", 1), ("final", 5)):
        tables.append(
            {
                "label": label,
                "header_row": 5,
                "start_row": 5,
                "end_row": 6,
                "start_column": start_column,
                "end_column": start_column + 2,
                "mapping_rules": _rules(derived_id=True),
            }
        )
    configure_migration_run(
        actor=actor,
        run=run,
        workbook_content=content,
        sheet_configurations=[
            {
                "sheet_name": "Synthetic Side By Side",
                "disposition": "included",
                "tables": tables,
            }
        ],
    )

    first = dry_run_migration(actor=actor, run=run, workbook_content=content)
    source_ids = list(
        first.row_previews.order_by("table_plan__label").values_list(
            "canonical_payload__source_result_id", flat=True
        )
    )

    assert len(source_ids) == 2
    assert len(set(source_ids)) == 2
    assert all(source_id.startswith("legacy:") for source_id in source_ids)


def test_dry_run_classifies_all_outcomes_without_results_side_effects() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Dry Run")
    actor = _actor(organization)
    seed_mapping = _identity_mapping(organization, actor)
    ingest_manual_rows(
        actor=actor,
        organization=organization,
        mapping_template=seed_mapping,
        operator_key="synthetic-legacy-seed",
        rows=[
            _canonical_row("existing-same", "Synthetic Same", "10.5"),
            _canonical_row("existing-conflict", "Synthetic Conflict", "11.5"),
        ],
    )
    rows = [
        ["new-result", "Synthetic New", "12.34"],
        ["invalid-result", "Synthetic Invalid", "not-a-number"],
        ["existing-same", "Synthetic Same", "10.5"],
        ["existing-conflict", "Synthetic Conflict", "99.9"],
        [None, None, None],
    ]
    content = _xlsx(
        [("Synthetic Results", "visible")],
        rows_by_sheet={"Synthetic Results": rows},
    )
    run = _inventory_run(content, organization, actor)
    configure_migration_run(
        actor=actor,
        run=run,
        workbook_content=content,
        sheet_configurations=_configuration(included_name="Synthetic Results", end_row=10),
    )
    before = {
        "artifacts": SourceArtifact.objects.count(),
        "runs": IngestionRun.objects.count(),
        "staged": StagedResult.objects.count(),
        "published": PublishedSourceResult.objects.count(),
        "cases": ReconciliationCase.objects.count(),
    }

    completed = dry_run_migration(actor=actor, run=run, workbook_content=content)

    assert list(completed.row_previews.values_list("source_row_number", "classification")) == [
        (6, "publishable"),
        (7, "quarantined"),
        (8, "duplicate"),
        (9, "conflict"),
        (10, "ignored"),
    ]
    assert completed.state == LegacyMigrationRun.State.DRY_RUN_COMPLETE
    assert (
        completed.publishable_rows,
        completed.quarantined_rows,
        completed.duplicate_rows,
        completed.conflict_rows,
        completed.ignored_rows,
    ) == (1, 1, 1, 1, 1)
    assert completed.dry_run_manifest_digest
    assert {
        "artifacts": SourceArtifact.objects.count(),
        "runs": IngestionRun.objects.count(),
        "staged": StagedResult.objects.count(),
        "published": PublishedSourceResult.objects.count(),
        "cases": ReconciliationCase.objects.count(),
    } == before
    manifest = LegacyMigrationManifest.objects.get(
        run=completed, kind=LegacyMigrationManifest.Kind.DRY_RUN
    )
    assert manifest.digest == completed.dry_run_manifest_digest
    assert "Synthetic New" not in str(manifest.payload)


def test_formula_in_any_mapped_cell_rolls_back_entire_dry_run() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Formula Rollback")
    actor = _actor(organization)
    content = _xlsx(
        [("Synthetic Formula", "visible")],
        rows_by_sheet={
            "Synthetic Formula": [
                ["valid-first", "Synthetic First", "12.3"],
                ["formula-second", "Synthetic Formula", "=10+2"],
            ]
        },
    )
    run = _inventory_run(content, organization, actor)
    configure_migration_run(
        actor=actor,
        run=run,
        workbook_content=content,
        sheet_configurations=_configuration(included_name="Synthetic Formula", end_row=7),
    )

    with pytest.raises(ValidationError, match="formula.*mapped cell"):
        dry_run_migration(actor=actor, run=run, workbook_content=content)

    assert LegacyMigrationRowPreview.objects.filter(run=run).count() == 0
    assert LegacyMigrationManifest.objects.filter(run=run).count() == 0
    run.refresh_from_db()
    assert run.state == LegacyMigrationRun.State.CONFIGURED
    assert run.dry_run_manifest_digest == ""


def test_dry_run_rejects_source_drift_before_persisting_preview_evidence() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Source Drift")
    actor = _actor(organization)
    original = _xlsx(
        [("Synthetic Drift", "visible")],
        rows_by_sheet={"Synthetic Drift": [["original", "Synthetic Original", "12.3"]]},
    )
    changed = _xlsx(
        [("Synthetic Drift", "visible")],
        rows_by_sheet={"Synthetic Drift": [["changed", "Synthetic Changed", "12.4"]]},
    )
    run = _inventory_run(original, organization, actor)
    configure_migration_run(
        actor=actor,
        run=run,
        workbook_content=original,
        sheet_configurations=_configuration(included_name="Synthetic Drift", end_row=6),
    )

    with pytest.raises(ValidationError, match="source artifact digest"):
        dry_run_migration(actor=actor, run=run, workbook_content=changed)

    assert LegacyMigrationRowPreview.objects.filter(run=run).count() == 0
    run.refresh_from_db()
    assert run.state == LegacyMigrationRun.State.CONFIGURED
