from __future__ import annotations

import hashlib
import json
import math
import time
from collections.abc import Callable
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from django.db import connection
from django.test import Client
from django.urls import reverse
from openpyxl import Workbook

from mnemex.accounts.models import Account, PrivilegedRoleAssignment
from mnemex.legacy_migration.inspection import inspect_legacy_workbook
from mnemex.legacy_migration.models import (
    LegacyMigrationCheckpoint,
    LegacyMigrationJob,
    LegacyMigrationManifest,
    LegacyMigrationRun,
    LegacySourceInventory,
)
from mnemex.legacy_migration.services import (
    MAX_APPLY_CHUNK_ROWS,
    apply_migration,
    approve_migration,
    configure_migration_run,
    dry_run_migration,
)
from mnemex.partners.models import PartnerOrganization
from mnemex.results.models import IngestionRun, PublishedSourceResult, SourceArtifact
from tests.mfa_helpers import enroll_account_mfa, force_login_with_fresh_mfa

pytestmark = pytest.mark.django_db


def _account(
    organization: PartnerOrganization,
    *,
    label: str,
    role: PrivilegedRoleAssignment.Role,
) -> Account:
    account = Account.objects.create_user(
        email=f"{label}-{organization.pk}@mnemex.example.invalid",
        password="synthetic-password-123",
    )
    enroll_account_mfa(account)
    PrivilegedRoleAssignment.objects.create(
        account=account,
        role=role,
        scope=PrivilegedRoleAssignment.Scope.ORGANIZATION,
        organization=organization,
        assigned_by=account,
    )
    return account


def _one_sheet_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic Private Results"
    sheet.append(["Result ID", "Competitor", "Score"])
    sheet.append(["synthetic-private-1", "Synthetic Private One", "10.1"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _run(
    organization: PartnerOrganization,
    actor: Account,
    *,
    source_label: str,
    content: bytes | None = None,
) -> LegacyMigrationRun:
    content = content or _one_sheet_workbook()
    digest = hashlib.sha256(content).hexdigest()
    artifact = SourceArtifact.objects.create(
        organization=organization,
        kind=SourceArtifact.Kind.SPREADSHEET,
        digest=digest,
        original_name=f"{source_label}.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        byte_size=len(content),
        object_reference=f"synthetic/private/{source_label}.xlsx",
        uploaded_by=actor,
    )
    inventory = LegacySourceInventory.objects.create(
        organization=organization,
        artifact=artifact,
        source_key=source_label,
        data_rights_reference="synthetic-test-authorization",
        parser_version="legacy-xlsx-inspector-v1",
        manifest_digest=hashlib.sha256(source_label.encode()).hexdigest(),
        created_by=actor,
    )
    return LegacyMigrationRun.objects.create(
        organization=organization,
        inventory=inventory,
        run_version=1,
        ruleset_version="legacy-mapping-v1",
        created_by=actor,
    )


def _reviewer(organization: PartnerOrganization, *, label: str) -> Account:
    return _account(
        organization,
        label=label,
        role=PrivilegedRoleAssignment.Role.EXPORT_REVIEWER,
    )


def _with_extra_archive_entries(content: bytes, *, count: int) -> bytes:
    stream = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(stream, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry, source.read(entry.filename))
        for index in range(count):
            target.writestr(
                f"synthetic-scale-metadata/entry-{index:03}.txt",
                f"synthetic-scale-entry-{index:03}",
            )
    return stream.getvalue()


def _layout_family_workbook(sheet_count: int) -> tuple[bytes, list[str]]:
    workbook = Workbook()
    workbook.remove(workbook.active)
    names: list[str] = []
    for index in range(sheet_count):
        suffix = " " if index == sheet_count - 1 else ""
        name = f"Synthetic {index + 1:02d}{suffix}"
        names.append(name)
        sheet = workbook.create_sheet(title=name)
        if index == 1:
            sheet.sheet_state = "hidden"
        sheet.merge_cells("A1:C1")
        sheet["A1"] = f"Synthetic scale layout {index + 1:02d}"
        for column, header in enumerate(
            (
                "Qual Result ID",
                "Qual Competitor",
                "Qual Score",
                "Qual Place",
                "Final Result ID",
                "Final Competitor",
                "Final Score",
            ),
            start=1,
        ):
            sheet.cell(row=5, column=column, value=header)
        for column, header in enumerate(
            ("Late Result ID", "Late Competitor", "Late Score"), start=1
        ):
            sheet.cell(row=12, column=column, value=header)
        sheet["H20"] = "=1+1"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return _with_extra_archive_entries(stream.getvalue(), count=105), names


def _large_workbook(row_count: int) -> bytes:
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = "Synthetic Scale Results"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Synthetic scale results only"
    sheet.cell(row=6, column=1, value="Result ID")
    sheet.cell(row=6, column=2, value="Competitor")
    sheet.cell(row=6, column=3, value="Score")
    for index in range(1, row_count + 1):
        row_number = index + 6
        sheet.cell(row=row_number, column=1, value=f"synthetic-scale-{index:05d}")
        sheet.cell(row=row_number, column=2, value=f"Synthetic Scale Person {index:05d}")
        sheet.cell(row=row_number, column=3, value=f"{10 + (index / 10_000):.4f}")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _rules() -> dict[str, dict[str, object]]:
    return {
        "source_result_id": {"kind": "column", "column": "Result ID"},
        "source_revision": {"kind": "constant", "value": 1},
        "source_event_id": {"kind": "constant", "value": "synthetic-scale-event-2026"},
        "event_name": {"kind": "constant", "value": "Synthetic Scale Show"},
        "result_date": {"kind": "constant", "value": "2026-08-01"},
        "competitor_name": {"kind": "column", "column": "Competitor"},
        "discipline": {"kind": "constant", "value": "UNDERHAND"},
        "score_type": {"kind": "constant", "value": "time"},
        "score": {"kind": "column", "column": "Score"},
    }


def _configured_scale_run(
    *,
    row_count: int,
    organization_label: str,
) -> tuple[LegacyMigrationRun, Account, Account, bytes]:
    organization = PartnerOrganization.objects.create(name=organization_label)
    manager = _account(
        organization,
        label=f"{organization_label.lower().replace(' ', '-')}-manager",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    reviewer = _reviewer(
        organization,
        label=f"{organization_label.lower().replace(' ', '-')}-reviewer",
    )
    content = _large_workbook(row_count)
    run = _run(
        organization,
        manager,
        source_label=f"synthetic-scale-{row_count}",
        content=content,
    )
    configure_migration_run(
        actor=manager,
        run=run,
        workbook_content=content,
        sheet_configurations=[
            {
                "sheet_name": "Synthetic Scale Results",
                "disposition": "included",
                "tables": [
                    {
                        "label": "synthetic-scale-table",
                        "header_row": 6,
                        "start_row": 6,
                        "end_row": row_count + 6,
                        "start_column": 1,
                        "end_column": 3,
                        "mapping_rules": _rules(),
                    }
                ],
            }
        ],
    )
    dry_run_migration(actor=manager, run=run, workbook_content=content)
    run.refresh_from_db()
    approve_migration(
        actor=reviewer,
        run=run,
        manifest_digest=run.dry_run_manifest_digest,
        rationale="Synthetic scale evidence reviewed.",
    )
    run.refresh_from_db()
    return run, manager, reviewer, content


class _QueryCounter:
    def __init__(self) -> None:
        self.count = 0

    def __call__(
        self,
        execute: object,
        sql: str,
        params: object,
        many: bool,
        context: object,
    ) -> object:
        self.count += 1
        return execute(sql, params, many, context)  # type: ignore[operator]


def _manifest_digest(payload: object) -> str:
    encoded = json.dumps(
        payload,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode()
    return hashlib.sha256(encoded).hexdigest()


@pytest.mark.parametrize("sheet_count", [28, 38, 42])
def test_layout_family_discovery_is_complete_and_deterministic(sheet_count: int) -> None:
    content, expected_names = _layout_family_workbook(sheet_count)

    first = inspect_legacy_workbook(content)
    replay = inspect_legacy_workbook(content)

    assert first == replay
    assert first.artifact_digest == hashlib.sha256(content).hexdigest()
    assert first.archive_entries > 100
    assert len(first.sheets) == sheet_count
    assert [sheet.name for sheet in first.sheets] == expected_names
    assert first.sheets[1].visibility == "hidden"
    assert first.sheets[-1].name.endswith(" ")
    for sheet in first.sheets:
        assert sheet.merged_ranges == ("A1:C1",)
        assert {5, 12}.issubset(sheet.header_candidates)
        assert sheet.formula_cells == ("H20",)
        assert sheet.max_row == 20
        assert sheet.max_column == 8
        assert len(sheet.metadata_digest) == 64


def test_more_than_two_thousand_rows_apply_in_bounded_replay_safe_partitions(
    record_property: Callable[[str, object], None],
) -> None:
    row_count = 2_001
    run, manager, _, content = _configured_scale_run(
        row_count=row_count,
        organization_label="Synthetic Large Scale Tenant",
    )
    dry_run_manifest = LegacyMigrationManifest.objects.get(
        run=run,
        kind=LegacyMigrationManifest.Kind.DRY_RUN,
    )
    original_digests = {
        "source": run.source_state_digest,
        "plan": run.plan_digest,
        "dry_run": run.dry_run_manifest_digest,
    }
    assert run.row_previews.count() == row_count
    assert dry_run_manifest.digest == _manifest_digest(dry_run_manifest.payload)
    assert len(dry_run_manifest.payload["rows"]) == row_count
    assert {row["classification"] for row in dry_run_manifest.payload["rows"]} == {"publishable"}
    assert {
        "total_rows": run.total_rows,
        "candidate_rows": run.candidate_rows,
        "publishable_rows": run.publishable_rows,
        "duplicate_rows": run.duplicate_rows,
        "conflict_rows": run.conflict_rows,
        "quarantined_rows": run.quarantined_rows,
        "ignored_rows": run.ignored_rows,
    } == {
        "total_rows": row_count,
        "candidate_rows": row_count,
        "publishable_rows": row_count,
        "duplicate_rows": 0,
        "conflict_rows": 0,
        "quarantined_rows": 0,
        "ignored_rows": 0,
    }

    query_counter = _QueryCounter()
    started = time.perf_counter()
    with connection.execute_wrapper(query_counter):
        completed = apply_migration(
            actor=manager,
            run=run,
            workbook_content=content,
            approved_manifest_digest=run.dry_run_manifest_digest,
        )
    elapsed = time.perf_counter() - started
    record_property("u5_scale_backend", connection.vendor)
    record_property("u5_scale_rows", row_count)
    record_property("u5_scale_apply_queries", query_counter.count)
    record_property("u5_scale_apply_elapsed_seconds", round(elapsed, 6))
    print(
        "U5 integration observation: "
        f"backend={connection.vendor} rows={row_count} "
        f"apply_queries={query_counter.count} elapsed_seconds={elapsed:.6f}"
    )

    checkpoints = list(completed.checkpoints.order_by("sequence"))
    reconciliation = LegacyMigrationManifest.objects.get(
        run=completed,
        kind=LegacyMigrationManifest.Kind.RECONCILIATION,
    )
    ingestion_runs = list(
        IngestionRun.objects.filter(legacy_migration_checkpoint__run=completed).order_by(
            "legacy_migration_checkpoint__sequence"
        )
    )
    assert completed.applied_rows == row_count
    assert len(checkpoints) == math.ceil(row_count / MAX_APPLY_CHUNK_ROWS)
    assert len(ingestion_runs) == len(checkpoints)
    assert all(1 <= checkpoint.total_rows <= MAX_APPLY_CHUNK_ROWS for checkpoint in checkpoints)
    assert all(1 <= ingestion.total_rows <= MAX_APPLY_CHUNK_ROWS for ingestion in ingestion_runs)
    assert sum(checkpoint.total_rows for checkpoint in checkpoints) == row_count
    assert sum(checkpoint.published_rows for checkpoint in checkpoints) == row_count
    assert (
        PublishedSourceResult.objects.filter(organization=completed.organization).count()
        == row_count
    )
    assert reconciliation.digest == _manifest_digest(reconciliation.payload)
    assert reconciliation.payload["totals"] == {
        "discovered_sheets": 1,
        "included_sheets": 1,
        "ignored_sheets": 0,
        "preview_rows": row_count,
        "candidate_rows": row_count,
        "ignored_rows": 0,
        "published_rows": row_count,
        "quarantined_rows": 0,
        "duplicate_rows": 0,
        "conflict_rows": 0,
    }
    checkpoint_digests = [checkpoint.checkpoint_digest for checkpoint in checkpoints]
    counts_before_replay = (
        LegacyMigrationCheckpoint.objects.filter(run=completed).count(),
        IngestionRun.objects.filter(legacy_migration_checkpoint__run=completed).count(),
        PublishedSourceResult.objects.filter(organization=completed.organization).count(),
    )

    replay = apply_migration(
        actor=manager,
        run=completed,
        workbook_content=content,
        approved_manifest_digest=run.dry_run_manifest_digest,
    )
    replay_reconciliation = LegacyMigrationManifest.objects.get(
        run=replay,
        kind=LegacyMigrationManifest.Kind.RECONCILIATION,
    )

    assert replay.source_state_digest == original_digests["source"]
    assert replay.plan_digest == original_digests["plan"]
    assert replay.dry_run_manifest_digest == original_digests["dry_run"]
    assert (
        list(replay.checkpoints.order_by("sequence").values_list("checkpoint_digest", flat=True))
        == checkpoint_digests
    )
    assert replay_reconciliation.digest == reconciliation.digest
    assert (
        LegacyMigrationCheckpoint.objects.filter(run=replay).count(),
        IngestionRun.objects.filter(legacy_migration_checkpoint__run=replay).count(),
        PublishedSourceResult.objects.filter(organization=replay.organization).count(),
    ) == counts_before_replay
    assert query_counter.count > 0


@pytest.mark.django_db(transaction=True)
def test_postgresql_committed_chunk_survives_failure_and_connection_boundary(
    record_property: Callable[[str, object], None],
) -> None:
    if connection.vendor != "postgresql":
        pytest.skip("requires the disposable PostgreSQL integration database")

    row_count = 501
    run, manager, _, content = _configured_scale_run(
        row_count=row_count,
        organization_label="Synthetic PostgreSQL Crash Tenant",
    )
    approved_digest = run.dry_run_manifest_digest
    started = time.perf_counter()
    with pytest.raises(RuntimeError, match="synthetic injected failure after committed chunk"):
        apply_migration(
            actor=manager,
            run=run,
            workbook_content=content,
            approved_manifest_digest=approved_digest,
            fail_after_committed_chunks=1,
        )
    interrupted_elapsed = time.perf_counter() - started

    assert not connection.in_atomic_block
    connection.close()
    connection.connect()

    run = LegacyMigrationRun.objects.get(pk=run.pk)
    committed = list(run.checkpoints.order_by("sequence"))
    # The durable worker keeps transient failures resumable instead of
    # collapsing them into a terminal run failure. The compatibility caller
    # may reclaim this retry-wait job immediately after the connection
    # boundary below.
    assert run.state == LegacyMigrationRun.State.APPLYING
    assert run.jobs.get(job_kind=LegacyMigrationJob.Kind.APPLY).status == (
        LegacyMigrationJob.Status.RETRY_WAIT
    )
    assert len(committed) == 1
    assert committed[0].status == LegacyMigrationCheckpoint.Status.COMMITTED
    assert committed[0].total_rows == MAX_APPLY_CHUNK_ROWS
    assert PublishedSourceResult.objects.filter(organization=run.organization).count() == (
        MAX_APPLY_CHUNK_ROWS
    )
    first_checkpoint_digest = committed[0].checkpoint_digest

    resume_counter = _QueryCounter()
    resumed_started = time.perf_counter()
    with connection.execute_wrapper(resume_counter):
        completed = apply_migration(
            actor=manager,
            run=run,
            workbook_content=content,
            approved_manifest_digest=approved_digest,
        )
    resumed_elapsed = time.perf_counter() - resumed_started
    record_property("u5_pg_rows", row_count)
    record_property("u5_pg_interrupted_elapsed_seconds", round(interrupted_elapsed, 6))
    record_property("u5_pg_resume_queries", resume_counter.count)
    record_property("u5_pg_resume_elapsed_seconds", round(resumed_elapsed, 6))
    print(
        "U5 PostgreSQL crash/resume observation: "
        f"rows={row_count} interrupted_seconds={interrupted_elapsed:.6f} "
        f"resume_queries={resume_counter.count} resume_seconds={resumed_elapsed:.6f}"
    )

    checkpoint_totals = list(
        completed.checkpoints.order_by("sequence").values_list("total_rows", flat=True)
    )
    assert completed.state == LegacyMigrationRun.State.COMPLETED
    assert checkpoint_totals == [250, 250, 1]
    assert completed.checkpoints.get(sequence=1).checkpoint_digest == first_checkpoint_digest
    assert (
        PublishedSourceResult.objects.filter(organization=completed.organization).count()
        == row_count
    )
    assert resume_counter.count > 0

    reconciliation = LegacyMigrationManifest.objects.get(
        run=completed,
        kind=LegacyMigrationManifest.Kind.RECONCILIATION,
    )
    reconciliation_digest = reconciliation.digest
    connection.close()
    connection.connect()
    replay = apply_migration(
        actor=manager,
        run=LegacyMigrationRun.objects.get(pk=completed.pk),
        workbook_content=content,
        approved_manifest_digest=approved_digest,
    )
    assert replay.checkpoints.count() == 3
    assert (
        LegacyMigrationManifest.objects.get(
            run=replay,
            kind=LegacyMigrationManifest.Kind.RECONCILIATION,
        ).digest
        == reconciliation_digest
    )


def test_cross_tenant_run_identifiers_are_non_enumerable() -> None:
    visible = PartnerOrganization.objects.create(name="Synthetic Visible Scale Tenant")
    hidden = PartnerOrganization.objects.create(name="Synthetic Hidden Scale Tenant")
    visible_manager = _account(
        visible,
        label="synthetic-visible-scale-manager",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    hidden_manager = _account(
        hidden,
        label="synthetic-hidden-scale-manager",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    _run(visible, visible_manager, source_label="synthetic-visible-scale")
    hidden_run = _run(hidden, hidden_manager, source_label="synthetic-hidden-scale")
    client = Client()
    force_login_with_fresh_mfa(client, visible_manager)

    dashboard = client.get(reverse("legacy_migration:dashboard"))
    dashboard_text = dashboard.content.decode()

    assert dashboard.status_code == 200
    assert str(hidden_run.pk) not in dashboard_text
    assert "Synthetic Hidden Scale Tenant" not in dashboard_text
    assert "synthetic-hidden-scale" not in dashboard_text
    for route_name in ("run-detail", "dry-run", "report-csv"):
        response = client.get(reverse(f"legacy_migration:{route_name}", args=[hidden_run.pk]))
        assert response.status_code == 404
