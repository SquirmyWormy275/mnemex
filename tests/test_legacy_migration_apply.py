from __future__ import annotations

import hashlib
from io import BytesIO

import pytest
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import models
from openpyxl import Workbook

from mnemex.accounts.models import Account, PrivilegedRoleAssignment
from mnemex.legacy_migration.models import (
    LegacyMigrationCheckpoint,
    LegacyMigrationDecisionRevision,
    LegacyMigrationJob,
    LegacyMigrationManifest,
    LegacyMigrationRowPreview,
    LegacyMigrationRun,
    LegacySourceInventory,
    LegacyTablePlan,
)
from mnemex.legacy_migration.services import (
    apply_migration,
    approve_migration,
    configure_migration_run,
    dry_run_migration,
    withdraw_migration,
)
from mnemex.partners.models import PartnerOrganization
from mnemex.results.models import (
    IngestionRun,
    MappingTemplate,
    PublishedSourceResult,
    ReconciliationCase,
    SourceArtifact,
    StagedResult,
)
from mnemex.results.services import ingest_manual_rows, ingest_preparsed_rows
from tests.mfa_helpers import enroll_account_mfa

pytestmark = pytest.mark.django_db


def _account(
    organization: PartnerOrganization,
    *,
    label: str,
    roles: tuple[PrivilegedRoleAssignment.Role, ...],
    mfa: bool = True,
) -> Account:
    actor = Account.objects.create_user(
        email=f"{label}-{organization.pk}@mnemex.example.invalid",
        password="synthetic-password-123",
    )
    if mfa:
        enroll_account_mfa(actor)
    for role in roles:
        PrivilegedRoleAssignment.objects.create(
            account=actor,
            role=role,
            scope=PrivilegedRoleAssignment.Scope.ORGANIZATION,
            organization=organization,
            assigned_by=actor,
        )
    return actor


def _xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic Results"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Synthetic legacy table"
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append(["Result ID", "Competitor", "Score"])
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _rules() -> dict[str, dict[str, object]]:
    return {
        "source_result_id": {"kind": "column", "column": "Result ID"},
        "source_revision": {"kind": "constant", "value": 1},
        "source_event_id": {"kind": "constant", "value": "synthetic-event-2026"},
        "event_name": {"kind": "constant", "value": "Synthetic Legacy Show"},
        "result_date": {"kind": "constant", "value": "2026-07-18"},
        "competitor_name": {"kind": "column", "column": "Competitor"},
        "discipline": {"kind": "constant", "value": "UNDERHAND"},
        "score_type": {"kind": "constant", "value": "time"},
        "score": {"kind": "column", "column": "Score"},
    }


def _canonical_row(result_id: str, competitor: str, score: str) -> dict[str, object]:
    return {
        "source_result_id": result_id,
        "source_revision": 1,
        "source_event_id": "synthetic-event-2026",
        "event_name": "Synthetic Legacy Show",
        "result_date": "2026-07-18",
        "competitor_name": competitor,
        "discipline": "UNDERHAND",
        "score_type": "time",
        "score": score,
    }


def _mapping(
    organization: PartnerOrganization,
    actor: Account,
    *,
    name: str = "Synthetic identity",
) -> MappingTemplate:
    fields = tuple(_canonical_row("id", "Synthetic", "1").keys())
    return MappingTemplate.objects.create(
        organization=organization,
        name=name,
        version=1,
        field_map={field: field for field in fields},
        created_by=actor,
    )


def _dry_run(
    organization: PartnerOrganization,
    manager: Account,
    rows: list[list[object]],
) -> tuple[LegacyMigrationRun, bytes]:
    content = _xlsx(rows)
    digest = hashlib.sha256(content).hexdigest()
    artifact = SourceArtifact.objects.create(
        organization=organization,
        kind=SourceArtifact.Kind.SPREADSHEET,
        digest=digest,
        original_name="synthetic-legacy.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        byte_size=len(content),
        object_reference="synthetic/private/legacy.xlsx",
        uploaded_by=manager,
    )
    inventory = LegacySourceInventory.objects.create(
        organization=organization,
        artifact=artifact,
        source_key=f"synthetic-source-{digest[:12]}",
        data_rights_reference="synthetic-test-authorization",
        parser_version="legacy-xlsx-inspector-v1",
        manifest_digest=hashlib.sha256(b"synthetic-inventory").hexdigest(),
        created_by=manager,
    )
    run = LegacyMigrationRun.objects.create(
        organization=organization,
        inventory=inventory,
        run_version=1,
        ruleset_version="legacy-mapping-v1",
        created_by=manager,
    )
    configure_migration_run(
        actor=manager,
        run=run,
        workbook_content=content,
        sheet_configurations=[
            {
                "sheet_name": "Synthetic Results",
                "disposition": "included",
                "tables": [
                    {
                        "label": "qualifying",
                        "header_row": 5,
                        "start_row": 5,
                        "end_row": 5 + len(rows),
                        "start_column": 1,
                        "end_column": 3,
                        "mapping_rules": _rules(),
                    }
                ],
            }
        ],
    )
    return dry_run_migration(actor=manager, run=run, workbook_content=content), content


def _roles() -> tuple[PartnerOrganization, Account, Account, LegacyMigrationRun, bytes]:
    organization = PartnerOrganization.objects.create(name="Synthetic U3 Organization")
    manager = _account(
        organization,
        label="manager",
        roles=(PrivilegedRoleAssignment.Role.RESULTS_MANAGER,),
    )
    reviewer = _account(
        organization,
        label="reviewer",
        roles=(PrivilegedRoleAssignment.Role.EXPORT_REVIEWER,),
    )
    run, content = _dry_run(
        organization,
        manager,
        [
            ["synthetic-1", "Synthetic One", "10.1"],
            ["synthetic-2", "Synthetic Two", "10.2"],
            ["synthetic-3", "Synthetic Three", "10.3"],
            ["synthetic-4", "Synthetic Four", "10.4"],
            ["synthetic-5", "Synthetic Five", "10.5"],
        ],
    )
    return organization, manager, reviewer, run, content


def test_partitioned_preparsed_ingestion_is_bounded_and_replay_safe() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Partition Intake")
    manager = _account(
        organization,
        label="partition-manager",
        roles=(PrivilegedRoleAssignment.Role.RESULTS_MANAGER,),
    )
    mapping = _mapping(organization, manager)
    artifact_digest = hashlib.sha256(b"synthetic-artifact").hexdigest()
    common = {
        "actor": manager,
        "organization": organization,
        "mapping_template": mapping,
        "rows": [_canonical_row("partition-1", "Synthetic Partition", "12.3")],
        "artifact_kind": SourceArtifact.Kind.SPREADSHEET,
        "artifact_digest": artifact_digest,
        "artifact_name": "synthetic.xlsx",
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "byte_size": 20,
        "object_reference": "synthetic/private/source.xlsx",
        "source_sheet_name": "Synthetic Results",
        "source_row_numbers": [6],
    }

    first = ingest_preparsed_rows(
        **common,
        operator_key="synthetic-partition-one",
        source_partition_key="synthetic-partition-one",
    )
    replay = ingest_preparsed_rows(
        **common,
        operator_key="synthetic-partition-one",
        source_partition_key="synthetic-partition-one",
    )
    second = ingest_preparsed_rows(
        **{
            **common,
            "rows": [_canonical_row("partition-2", "Synthetic Second", "12.4")],
        },
        operator_key="synthetic-partition-two",
        source_partition_key="synthetic-partition-two",
    )
    with pytest.raises(ValueError, match="different payload"):
        ingest_preparsed_rows(
            **{
                **common,
                "rows": [_canonical_row("changed", "Synthetic Changed", "99.9")],
            },
            operator_key="synthetic-partition-one",
            source_partition_key="synthetic-partition-one",
        )

    assert replay.replayed is True
    assert replay.run.pk == first.run.pk
    assert second.run.pk != first.run.pk
    assert set(IngestionRun.objects.values_list("source_partition_key", flat=True)) == {
        "synthetic-partition-one",
        "synthetic-partition-two",
    }


def test_approval_requires_exact_digest_mfa_role_and_tenant() -> None:
    organization, manager, reviewer, run, _ = _roles()
    other = PartnerOrganization.objects.create(name="Synthetic Other Organization")
    wrong_tenant_reviewer = _account(
        other,
        label="wrong-tenant-reviewer",
        roles=(PrivilegedRoleAssignment.Role.EXPORT_REVIEWER,),
    )
    no_mfa_reviewer = _account(
        organization,
        label="no-mfa-reviewer",
        roles=(PrivilegedRoleAssignment.Role.EXPORT_REVIEWER,),
        mfa=False,
    )

    for actor in (manager, wrong_tenant_reviewer, no_mfa_reviewer):
        with pytest.raises(PermissionDenied, match="export reviewer"):
            approve_migration(
                actor=actor,
                run=run,
                manifest_digest=run.dry_run_manifest_digest,
                rationale="Synthetic review approval.",
            )
    with pytest.raises(ValidationError, match="exact dry-run manifest"):
        approve_migration(
            actor=reviewer,
            run=run,
            manifest_digest="0" * 64,
            rationale="Synthetic stale review.",
        )

    approved = approve_migration(
        actor=reviewer,
        run=run,
        manifest_digest=run.dry_run_manifest_digest,
        rationale="Synthetic evidence reviewed.",
    )

    assert approved.state == LegacyMigrationRun.State.APPROVED
    decision = LegacyMigrationDecisionRevision.objects.get(run=run)
    assert decision.decision == LegacyMigrationDecisionRevision.Decision.APPROVE
    assert decision.manifest_digest == run.dry_run_manifest_digest
    assert decision.sequence == 1


def test_apply_requires_approval_manager_role_exact_digest_and_source() -> None:
    _, manager, reviewer, run, content = _roles()

    with pytest.raises(ValidationError, match="approved"):
        apply_migration(
            actor=manager,
            run=run,
            workbook_content=content,
            approved_manifest_digest=run.dry_run_manifest_digest,
        )
    approve_migration(
        actor=reviewer,
        run=run,
        manifest_digest=run.dry_run_manifest_digest,
        rationale="Synthetic evidence reviewed.",
    )
    with pytest.raises(PermissionDenied, match="results manager"):
        apply_migration(
            actor=reviewer,
            run=run,
            workbook_content=b"not-a-workbook-and-not-the-approved-source",
            approved_manifest_digest=run.dry_run_manifest_digest,
        )
    with pytest.raises(ValidationError, match="approved manifest"):
        apply_migration(
            actor=manager,
            run=run,
            workbook_content=content,
            approved_manifest_digest="f" * 64,
        )
    with pytest.raises(ValidationError, match="source artifact"):
        apply_migration(
            actor=manager,
            run=run,
            workbook_content=_xlsx([["changed", "Synthetic Changed", "99"]]),
            approved_manifest_digest=run.dry_run_manifest_digest,
        )


def test_approved_manifest_rejects_appended_or_digest_tampered_preview_rows() -> None:
    _, manager, reviewer, run, content = _roles()
    approve_migration(
        actor=reviewer,
        run=run,
        manifest_digest=run.dry_run_manifest_digest,
        rationale="Synthetic evidence reviewed.",
    )
    original = run.row_previews.select_related("table_plan").order_by("source_row_number").first()
    assert original is not None
    injected_table = LegacyTablePlan(
        organization=run.organization,
        run=run,
        worksheet=original.table_plan.worksheet,
        label="injected-after-approval",
        header_row=original.table_plan.header_row,
        start_row=original.table_plan.start_row,
        end_row=original.table_plan.end_row,
        start_column=original.table_plan.start_column,
        end_column=original.table_plan.end_column,
        mapping_rules=original.table_plan.mapping_rules,
        mapping_digest=original.table_plan.mapping_digest,
        created_by=manager,
    )
    models.Model.save(injected_table, force_insert=True)
    appended = LegacyMigrationRowPreview(
        organization=run.organization,
        run=run,
        table_plan=injected_table,
        source_sheet_index=original.source_sheet_index,
        source_row_number=original.source_row_number,
        source_column_start=original.source_column_start,
        source_column_end=original.source_column_end,
        canonical_payload=original.canonical_payload,
        payload_digest=original.payload_digest,
        classification=original.classification,
        validation_errors=original.validation_errors,
        source_state_digest=original.source_state_digest,
    )

    with pytest.raises(PermissionDenied, match="closed"):
        appended.save()

    # Simulate provenance inserted outside the application model guard. Apply
    # must still reconstruct and verify the exact approved preview set.
    appended.canonical_payload = {**original.canonical_payload, "score": "999"}
    models.Model.save(appended, force_insert=True)
    with pytest.raises(ValidationError, match="preview evidence"):
        apply_migration(
            actor=manager,
            run=run,
            workbook_content=content,
            approved_manifest_digest=run.dry_run_manifest_digest,
        )

    assert not IngestionRun.objects.exists()
    assert not LegacyMigrationCheckpoint.objects.exists()


def test_apply_commits_a_chunk_then_resumes_exactly_and_reconciles_every_row() -> None:
    _, manager, reviewer, run, content = _roles()
    approve_migration(
        actor=reviewer,
        run=run,
        manifest_digest=run.dry_run_manifest_digest,
        rationale="Synthetic evidence reviewed.",
    )

    with pytest.raises(RuntimeError, match="synthetic injected failure"):
        apply_migration(
            actor=manager,
            run=run,
            workbook_content=content,
            approved_manifest_digest=run.dry_run_manifest_digest,
            chunk_size=2,
            fail_after_committed_chunks=1,
        )

    run.refresh_from_db()
    assert run.state == LegacyMigrationRun.State.APPLYING
    assert PublishedSourceResult.objects.filter(organization=run.organization).count() == 2
    assert (
        LegacyMigrationCheckpoint.objects.filter(
            run=run, status=LegacyMigrationCheckpoint.Status.COMMITTED
        ).count()
        == 1
    )
    assert LegacyMigrationJob.objects.get(run=run, job_kind="apply").status == "retry_wait"
    interrupted_job = LegacyMigrationJob.objects.get(run=run, job_kind="apply")
    assert interrupted_job.last_error_code == "worker_error"
    assert interrupted_job.last_error_message == ("temporary worker failure; retry is scheduled")

    completed = apply_migration(
        actor=manager,
        run=run,
        workbook_content=content,
        approved_manifest_digest=run.dry_run_manifest_digest,
        chunk_size=2,
    )
    checkpoint_ids = list(completed.checkpoints.values_list("checkpoint_id", flat=True))
    replay = apply_migration(
        actor=manager,
        run=completed,
        workbook_content=content,
        approved_manifest_digest=run.dry_run_manifest_digest,
        chunk_size=2,
    )

    assert replay.state == LegacyMigrationRun.State.COMPLETED
    assert list(replay.checkpoints.values_list("checkpoint_id", flat=True)) == checkpoint_ids
    assert replay.applied_rows == 5
    assert replay.checkpoints.count() == 3
    assert IngestionRun.objects.filter(legacy_migration_checkpoint__run=run).count() == 3
    assert PublishedSourceResult.objects.filter(organization=run.organization).count() == 5
    assert not PublishedSourceResult.objects.filter(
        organization=run.organization, person__isnull=False
    ).exists()
    assert (
        ReconciliationCase.objects.filter(
            staged_result__run__legacy_migration_checkpoint__run=run,
            case_type=ReconciliationCase.CaseType.IDENTITY_UNRESOLVED,
        ).count()
        == 5
    )

    reconciliation = LegacyMigrationManifest.objects.get(
        run=run, kind=LegacyMigrationManifest.Kind.RECONCILIATION
    )
    assert reconciliation.payload["totals"] == {
        "discovered_sheets": 1,
        "included_sheets": 1,
        "ignored_sheets": 0,
        "preview_rows": 5,
        "candidate_rows": 5,
        "ignored_rows": 0,
        "published_rows": 5,
        "quarantined_rows": 0,
        "duplicate_rows": 0,
        "conflict_rows": 0,
    }
    assert len(reconciliation.payload["rows"]) == 5


def test_resume_rejects_lower_level_checkpoint_ingestion_tamper() -> None:
    _, manager, reviewer, run, content = _roles()
    approve_migration(
        actor=reviewer,
        run=run,
        manifest_digest=run.dry_run_manifest_digest,
        rationale="Synthetic evidence reviewed.",
    )
    with pytest.raises(RuntimeError, match="synthetic injected failure"):
        apply_migration(
            actor=manager,
            run=run,
            workbook_content=content,
            approved_manifest_digest=run.dry_run_manifest_digest,
            chunk_size=2,
            fail_after_committed_chunks=1,
        )
    checkpoint = run.checkpoints.select_related("ingestion_run").get(sequence=1)
    assert checkpoint.ingestion_run is not None
    staged = checkpoint.ingestion_run.staged_results.order_by("row_number").first()
    assert staged is not None

    staged.payload_digest = "0" * 64
    with pytest.raises(PermissionDenied, match="immutable"):
        staged.save()
    with pytest.raises(PermissionDenied, match="immutable"):
        StagedResult.objects.filter(pk=staged.pk).update(payload_digest="0" * 64)

    models.Model.save(staged, update_fields=["payload_digest"])
    with pytest.raises(ValidationError, match="checkpoint ingestion evidence"):
        apply_migration(
            actor=manager,
            run=run,
            workbook_content=content,
            approved_manifest_digest=run.dry_run_manifest_digest,
            chunk_size=2,
        )

    assert not LegacyMigrationManifest.objects.filter(
        run=run, kind=LegacyMigrationManifest.Kind.RECONCILIATION
    ).exists()


def test_apply_reclassifies_new_duplicate_and_conflict_at_publication_time() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Apply Drift")
    manager = _account(
        organization,
        label="drift-manager",
        roles=(PrivilegedRoleAssignment.Role.RESULTS_MANAGER,),
    )
    reviewer = _account(
        organization,
        label="drift-reviewer",
        roles=(PrivilegedRoleAssignment.Role.EXPORT_REVIEWER,),
    )
    run, content = _dry_run(
        organization,
        manager,
        [
            ["late-duplicate", "Synthetic Duplicate", "10.1"],
            ["late-conflict", "Synthetic Conflict", "10.2"],
        ],
    )
    mapping = _mapping(organization, manager, name="Synthetic late state")
    ingest_manual_rows(
        actor=manager,
        organization=organization,
        mapping_template=mapping,
        operator_key="synthetic-late-state",
        rows=[
            _canonical_row("late-duplicate", "Synthetic Duplicate", "10.1"),
            _canonical_row("late-conflict", "Synthetic Conflict", "99.9"),
        ],
    )
    approve_migration(
        actor=reviewer,
        run=run,
        manifest_digest=run.dry_run_manifest_digest,
        rationale="Synthetic evidence reviewed before apply.",
    )

    apply_migration(
        actor=manager,
        run=run,
        workbook_content=content,
        approved_manifest_digest=run.dry_run_manifest_digest,
    )

    outcomes = set(
        StagedResult.objects.filter(run__legacy_migration_checkpoint__run=run).values_list(
            "outcome", flat=True
        )
    )
    reconciliation = LegacyMigrationManifest.objects.get(
        run=run, kind=LegacyMigrationManifest.Kind.RECONCILIATION
    )
    assert outcomes == {StagedResult.Outcome.DUPLICATE, StagedResult.Outcome.CONFLICT}
    assert reconciliation.payload["totals"]["duplicate_rows"] == 1
    assert reconciliation.payload["totals"]["conflict_rows"] == 1
    assert {row["dry_run_classification"] for row in reconciliation.payload["rows"]} == {
        "publishable"
    }
    assert {row["apply_outcome"] for row in reconciliation.payload["rows"]} == {
        "duplicate",
        "conflict",
    }


def test_withdrawal_is_append_only_retains_evidence_and_blocks_future_apply() -> None:
    _, manager, reviewer, run, content = _roles()
    approve_migration(
        actor=reviewer,
        run=run,
        manifest_digest=run.dry_run_manifest_digest,
        rationale="Synthetic evidence reviewed.",
    )
    completed = apply_migration(
        actor=manager,
        run=run,
        workbook_content=content,
        approved_manifest_digest=run.dry_run_manifest_digest,
        chunk_size=2,
    )
    reconciliation = LegacyMigrationManifest.objects.get(
        run=run, kind=LegacyMigrationManifest.Kind.RECONCILIATION
    )
    artifact_ids = set(SourceArtifact.objects.values_list("artifact_id", flat=True))
    published_ids = set(PublishedSourceResult.objects.values_list("published_result_id", flat=True))

    withdrawn = withdraw_migration(
        actor=reviewer,
        run=completed,
        manifest_digest=reconciliation.digest,
        rationale="Synthetic source authorization withdrawn.",
    )

    assert withdrawn.state == LegacyMigrationRun.State.WITHDRAWN
    assert set(SourceArtifact.objects.values_list("artifact_id", flat=True)) == artifact_ids
    assert (
        set(PublishedSourceResult.objects.values_list("published_result_id", flat=True))
        == published_ids
    )
    decisions = list(run.decision_revisions.order_by("sequence"))
    assert [decision.decision for decision in decisions] == ["approve", "withdraw"]
    assert decisions[1].predecessor == decisions[0]
    withdrawal = LegacyMigrationManifest.objects.get(
        run=run, kind=LegacyMigrationManifest.Kind.WITHDRAWAL
    )
    assert set(withdrawal.payload["published_result_ids"]) == {
        str(value) for value in published_ids
    }
    with pytest.raises(ValidationError, match="withdrawn"):
        apply_migration(
            actor=manager,
            run=withdrawn,
            workbook_content=content,
            approved_manifest_digest=run.dry_run_manifest_digest,
        )


def test_side_by_side_tables_can_checkpoint_the_same_source_row_range() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Side By Side")
    manager = _account(
        organization,
        label="side-manager",
        roles=(PrivilegedRoleAssignment.Role.RESULTS_MANAGER,),
    )
    reviewer = _account(
        organization,
        label="side-reviewer",
        roles=(PrivilegedRoleAssignment.Role.EXPORT_REVIEWER,),
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic Parallel"
    for start_column in (1, 5):
        for offset, header in enumerate(("Result ID", "Competitor", "Score")):
            sheet.cell(row=5, column=start_column + offset, value=header)
    for column, value in enumerate(("parallel-left", "Synthetic Left", "11.1"), start=1):
        sheet.cell(row=6, column=column, value=value)
    for column, value in enumerate(("parallel-right", "Synthetic Right", "11.2"), start=5):
        sheet.cell(row=6, column=column, value=value)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    content = stream.getvalue()
    digest = hashlib.sha256(content).hexdigest()
    artifact = SourceArtifact.objects.create(
        organization=organization,
        kind=SourceArtifact.Kind.SPREADSHEET,
        digest=digest,
        original_name="synthetic-parallel.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        byte_size=len(content),
        object_reference="synthetic/private/parallel.xlsx",
        uploaded_by=manager,
    )
    inventory = LegacySourceInventory.objects.create(
        organization=organization,
        artifact=artifact,
        source_key="synthetic-parallel-source",
        data_rights_reference="synthetic-test-authorization",
        parser_version="legacy-xlsx-inspector-v1",
        manifest_digest=hashlib.sha256(b"synthetic-parallel-inventory").hexdigest(),
        created_by=manager,
    )
    run = LegacyMigrationRun.objects.create(
        organization=organization,
        inventory=inventory,
        run_version=1,
        ruleset_version="legacy-mapping-v1",
        created_by=manager,
    )
    table_base = {
        "header_row": 5,
        "start_row": 5,
        "end_row": 6,
        "mapping_rules": _rules(),
    }
    configure_migration_run(
        actor=manager,
        run=run,
        workbook_content=content,
        sheet_configurations=[
            {
                "sheet_name": "Synthetic Parallel",
                "disposition": "included",
                "tables": [
                    {
                        **table_base,
                        "label": "left",
                        "start_column": 1,
                        "end_column": 3,
                    },
                    {
                        **table_base,
                        "label": "right",
                        "start_column": 5,
                        "end_column": 7,
                    },
                ],
            }
        ],
    )
    dry_run_migration(actor=manager, run=run, workbook_content=content)
    run.refresh_from_db()
    approve_migration(
        actor=reviewer,
        run=run,
        manifest_digest=run.dry_run_manifest_digest,
        rationale="Synthetic parallel evidence reviewed.",
    )

    apply_migration(
        actor=manager,
        run=run,
        workbook_content=content,
        approved_manifest_digest=run.dry_run_manifest_digest,
        chunk_size=1,
    )

    checkpoints = list(run.checkpoints.order_by("sequence"))
    assert len(checkpoints) == 2
    assert {
        (checkpoint.first_source_row, checkpoint.last_source_row) for checkpoint in checkpoints
    } == {(6, 6)}
    assert len({checkpoint.table_plan_id for checkpoint in checkpoints}) == 2
