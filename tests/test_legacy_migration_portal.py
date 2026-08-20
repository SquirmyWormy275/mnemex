from __future__ import annotations

import csv
import hashlib
import json
import uuid
from io import BytesIO, StringIO
from pathlib import Path

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from mnemex.accounts.models import Account, PrivilegedRoleAssignment
from mnemex.legacy_migration.models import LegacyMigrationJob, LegacyMigrationRun
from mnemex.legacy_migration.worker_jobs import run_apply_jobs_once
from mnemex.partners.models import PartnerOrganization
from mnemex.results.models import PublishedSourceResult, SourceArtifact
from tests.mfa_helpers import (
    enroll_account_mfa,
    force_login_with_fresh_mfa,
    verify_account_email,
)

pytestmark = pytest.mark.django_db


def _account(
    organization: PartnerOrganization,
    *,
    label: str,
    role: PrivilegedRoleAssignment.Role,
    mfa: bool = True,
) -> Account:
    account = Account.objects.create_user(
        email=f"{label}@mnemex.example.invalid", password="synthetic-password-123"
    )
    verify_account_email(account)
    if mfa:
        enroll_account_mfa(account)
    PrivilegedRoleAssignment.objects.create(
        account=account,
        role=role,
        scope=PrivilegedRoleAssignment.Scope.ORGANIZATION,
        organization=organization,
        assigned_by=account,
    )
    return account


def _xlsx() -> bytes:
    workbook = Workbook()
    results = workbook.active
    results.title = "Synthetic Results "
    results.merge_cells("A1:C1")
    results["A1"] = "Synthetic legacy table"
    results.append([])
    results.append([])
    results.append([])
    results.append(["Result ID", "Competitor", "Score"])
    results.append(["synthetic-1", "Synthetic Person One", "10.1"])
    summary = workbook.create_sheet("Synthetic Summary")
    summary.sheet_state = "hidden"
    summary.append(["Synthetic", "Only"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _configuration() -> str:
    rules = {
        "source_result_id": {"kind": "column", "column": "Result ID"},
        "source_revision": {"kind": "constant", "value": 1},
        "source_event_id": {"kind": "constant", "value": "synthetic-event-2026"},
        "event_name": {"kind": "constant", "value": "Synthetic Legacy Show"},
        "result_date": {"kind": "constant", "value": "2026-07-18"},
        "competitor_name": {"kind": "column", "column": "Competitor"},
        "discipline": {"kind": "constant", "value": "UNDERHAND"},
        "score_type": {"kind": "constant", "value": "time"},
        "score": {"kind": "column", "column": "Score"},
    }
    return json.dumps(
        [
            {
                "sheet_name": "Synthetic Results ",
                "disposition": "included",
                "tables": [
                    {
                        "label": "qualifying",
                        "header_row": 5,
                        "start_row": 5,
                        "end_row": 6,
                        "start_column": 1,
                        "end_column": 3,
                        "mapping_rules": rules,
                    }
                ],
            },
            {
                "sheet_name": "Synthetic Summary",
                "disposition": "ignored",
                "ignore_reason": "Synthetic summary is not row-level results.",
                "tables": [],
            },
        ]
    )


def _upload(
    client: Client,
    organization: PartnerOrganization,
    content: bytes,
    artifact_root: Path,
) -> LegacyMigrationRun:
    with override_settings(MNEMEX_PRIVATE_ARTIFACT_ROOT=artifact_root):
        response = client.post(
            reverse("legacy_migration:upload"),
            {
                "organization": str(organization.pk),
                "source_key": "synthetic-source-2026",
                "data_rights_reference": "synthetic-test-authorization",
                "spreadsheet": SimpleUploadedFile(
                    "synthetic-legacy.xlsx",
                    content,
                    content_type=(
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    ),
                ),
            },
        )
    assert response.status_code == 302
    return LegacyMigrationRun.objects.get()


def test_legacy_workspace_fails_closed_for_anonymous_missing_mfa_and_wrong_role() -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Access Show")
    anonymous = Client()
    response = anonymous.get(reverse("legacy_migration:dashboard"))
    assert response.status_code == 302
    assert response.url.startswith(reverse("login"))

    missing_mfa = Client()
    missing_mfa.force_login(
        _account(
            organization,
            label="legacy-no-mfa",
            role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
            mfa=False,
        )
    )
    missing = missing_mfa.get(reverse("legacy_migration:dashboard"))
    assert missing.status_code == 302
    assert missing.url.startswith(reverse("mfa_activate_totp"))

    ordinary = Account.objects.create_user(
        email="legacy-ordinary@mnemex.example.invalid",
        password="synthetic-password-123",
    )
    ordinary.mfa_enrolled_at = timezone.now()
    ordinary.save(update_fields=["mfa_enrolled_at"])
    wrong_role = Client()
    wrong_role.force_login(ordinary)
    assert wrong_role.get(reverse("legacy_migration:dashboard")).status_code == 403


def test_manager_upload_discovers_exact_structure_and_keeps_artifact_private(
    tmp_path: Path,
) -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Discovery Show")
    manager = _account(
        organization,
        label="legacy-discovery-manager",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    client = Client()
    force_login_with_fresh_mfa(client, manager)
    content = _xlsx()
    run = _upload(client, organization, content, tmp_path)

    with override_settings(MNEMEX_PRIVATE_ARTIFACT_ROOT=tmp_path):
        response = client.get(reverse("legacy_migration:run-detail", args=[run.pk]))

    assert response.status_code == 200
    body = response.content.decode()
    assert "Synthetic Results " in body
    assert "Synthetic Summary" in body
    assert "Hidden" in body
    assert "5" in body
    assert run.inventory.artifact.object_reference not in body
    assert "Synthetic Person One" not in body
    assert run.state == LegacyMigrationRun.State.DISCOVERED
    assert SourceArtifact.objects.get().digest == hashlib.sha256(content).hexdigest()


def test_invalid_workbook_cleans_unreferenced_private_object(tmp_path: Path) -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Invalid Workbook Show")
    manager = _account(
        organization,
        label="legacy-invalid-manager",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    client = Client()
    force_login_with_fresh_mfa(client, manager)

    with override_settings(MNEMEX_PRIVATE_ARTIFACT_ROOT=tmp_path):
        response = client.post(
            reverse("legacy_migration:upload"),
            {
                "organization": str(organization.pk),
                "source_key": "synthetic-invalid-source",
                "data_rights_reference": "synthetic-test-authorization",
                "spreadsheet": SimpleUploadedFile(
                    "synthetic-invalid.xlsx",
                    b"not an XLSX archive",
                    content_type=(
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    ),
                ),
            },
        )

    assert response.status_code == 200
    assert "not a valid workbook archive" in response.content.decode()
    assert SourceArtifact.objects.count() == 0
    assert not [path for path in tmp_path.rglob("*") if path.is_file()]


def test_tenant_scoping_hides_runs_and_denies_cross_tenant_detail(
    tmp_path: Path,
) -> None:
    organization_a = PartnerOrganization.objects.create(name="Visible Legacy Show")
    organization_b = PartnerOrganization.objects.create(name="Hidden Legacy Show")
    manager_a = _account(
        organization_a,
        label="legacy-manager-a",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    manager_b = _account(
        organization_b,
        label="legacy-manager-b",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    client_b = Client()
    force_login_with_fresh_mfa(client_b, manager_b)
    run_b = _upload(client_b, organization_b, _xlsx(), tmp_path)

    client_a = Client()
    force_login_with_fresh_mfa(client_a, manager_a)
    dashboard = client_a.get(reverse("legacy_migration:dashboard"))
    assert dashboard.status_code == 200
    assert "Hidden Legacy Show" not in dashboard.content.decode()
    assert client_a.get(reverse("legacy_migration:run-detail", args=[run_b.pk])).status_code == 404


def test_upload_nonexistent_and_unauthorized_organization_are_indistinguishable() -> None:
    managed = PartnerOrganization.objects.create(name="Synthetic upload oracle managed")
    unauthorized = PartnerOrganization.objects.create(name="Synthetic upload oracle hidden")
    manager = _account(
        managed,
        label="legacy-upload-oracle-manager",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    client = Client()
    force_login_with_fresh_mfa(client, manager)

    def post(organization_id: object):
        return client.post(
            reverse("legacy_migration:upload"),
            {
                "organization": str(organization_id),
                "source_key": "synthetic-oracle",
                "data_rights_reference": "synthetic-test-authorization",
                "spreadsheet": SimpleUploadedFile(
                    "synthetic-oracle.xlsx",
                    _xlsx(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    ),
                ),
            },
        )

    unauthorized_response = post(unauthorized.pk)
    nonexistent_response = post(uuid.uuid4())

    assert unauthorized_response.status_code == nonexistent_response.status_code == 200
    unauthorized_errors = unauthorized_response.context["form"].errors.as_json()
    nonexistent_errors = nonexistent_response.context["form"].errors.as_json()
    assert unauthorized_errors == nonexistent_errors


def test_manager_configures_and_dry_runs_without_publication(tmp_path: Path) -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Dry Run Show")
    manager = _account(
        organization,
        label="legacy-dry-run-manager",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    client = Client()
    force_login_with_fresh_mfa(client, manager)
    run = _upload(client, organization, _xlsx(), tmp_path)

    with override_settings(MNEMEX_PRIVATE_ARTIFACT_ROOT=tmp_path):
        configured = client.post(
            reverse("legacy_migration:configure", args=[run.pk]),
            {"configuration": _configuration()},
        )
        previewed = client.post(reverse("legacy_migration:dry-run", args=[run.pk]))

    assert configured.status_code == 302
    assert previewed.status_code == 302
    run.refresh_from_db()
    assert run.state == LegacyMigrationRun.State.DRY_RUN_COMPLETE
    assert run.publishable_rows == 1
    assert run.ignored_rows == 0
    assert PublishedSourceResult.objects.count() == 0
    detail = client.get(reverse("legacy_migration:run-detail", args=[run.pk]))
    body = detail.content.decode()
    assert "Five-way dry-run accounting" in body
    assert "publishable" in body.lower()
    assert "No result rows were published by this dry run" in body
    assert 'class="stacked-note"' in body
    assert "Synthetic Person One" not in body


def test_separate_reviewer_approval_apply_reconciliation_csv_and_withdrawal(
    tmp_path: Path,
) -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Approval Show")
    manager = _account(
        organization,
        label="legacy-apply-manager",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    reviewer = _account(
        organization,
        label="legacy-apply-reviewer",
        role=PrivilegedRoleAssignment.Role.EXPORT_REVIEWER,
    )
    manager_client = Client()
    force_login_with_fresh_mfa(manager_client, manager)
    run = _upload(manager_client, organization, _xlsx(), tmp_path)
    with override_settings(MNEMEX_PRIVATE_ARTIFACT_ROOT=tmp_path):
        assert (
            manager_client.post(
                reverse("legacy_migration:configure", args=[run.pk]),
                {"configuration": _configuration()},
            ).status_code
            == 302
        )
        assert (
            manager_client.post(reverse("legacy_migration:dry-run", args=[run.pk])).status_code
            == 302
        )
    run.refresh_from_db()

    assert (
        manager_client.post(
            reverse("legacy_migration:approve", args=[run.pk]),
            {
                "manifest_digest": run.dry_run_manifest_digest,
                "rationale": "Not my authority.",
            },
        ).status_code
        == 403
    )

    reviewer_client = Client()
    force_login_with_fresh_mfa(reviewer_client, reviewer)
    approved = reviewer_client.post(
        reverse("legacy_migration:approve", args=[run.pk]),
        {
            "manifest_digest": run.dry_run_manifest_digest,
            "rationale": "Synthetic reviewer approved the exact manifest.",
        },
    )
    assert approved.status_code == 302

    with override_settings(MNEMEX_PRIVATE_ARTIFACT_ROOT=tmp_path):
        stale_apply = manager_client.post(
            reverse("legacy_migration:apply", args=[run.pk]),
            {
                "manifest_digest": "0" * 64,
                "rationale": "Synthetic manager request.",
            },
        )
    assert stale_apply.status_code == 200
    assert "approved manifest digest" in stale_apply.content.decode().lower()
    run.refresh_from_db()
    assert run.state == LegacyMigrationRun.State.APPROVED
    assert PublishedSourceResult.objects.count() == 0

    with override_settings(MNEMEX_PRIVATE_ARTIFACT_ROOT=tmp_path):
        applied = manager_client.post(
            reverse("legacy_migration:apply", args=[run.pk]),
            {
                "manifest_digest": run.dry_run_manifest_digest,
                "rationale": "Synthetic manager request.",
            },
        )
    assert applied.status_code == 302
    run.refresh_from_db()
    queued_job = run.jobs.get(job_kind=LegacyMigrationJob.Kind.APPLY)
    assert run.state == LegacyMigrationRun.State.APPROVED
    assert queued_job.status == LegacyMigrationJob.Status.PENDING
    assert run.checkpoints.count() == 0
    assert PublishedSourceResult.objects.count() == 0

    worker_outcome = run_apply_jobs_once(
        artifact_root=tmp_path,
        owner="synthetic-portal-worker",
        limit=1,
    )
    assert worker_outcome.claimed == worker_outcome.completed == 1
    run.refresh_from_db()
    assert run.state == LegacyMigrationRun.State.COMPLETED
    assert run.checkpoints.count() == 1
    detail = manager_client.get(reverse("legacy_migration:run-detail", args=[run.pk]))
    assert "processed rows" in detail.content.decode().lower()

    report = reviewer_client.get(reverse("legacy_migration:report-csv", args=[run.pk]))
    assert report.status_code == 200
    rows = list(csv.DictReader(StringIO(report.content.decode())))
    assert rows[0]["source_row_number"] == "6"
    assert "competitor" not in report.content.decode().lower()
    assert "synthetic results" not in report.content.decode().lower()
    assert "object_reference" not in report.content.decode().lower()
    assert "source_result_id" not in report.content.decode().lower()

    reconciliation = run.manifests.get(kind="reconciliation")
    withdrawn = reviewer_client.post(
        reverse("legacy_migration:withdraw", args=[run.pk]),
        {
            "manifest_digest": reconciliation.digest,
            "rationale": "Synthetic withdrawal retains all immutable evidence.",
        },
    )
    assert withdrawn.status_code == 302
    run.refresh_from_db()
    assert run.state == LegacyMigrationRun.State.WITHDRAWN
    assert PublishedSourceResult.objects.count() == 1

    with override_settings(MNEMEX_PRIVATE_ARTIFACT_ROOT=tmp_path):
        reapplied = manager_client.post(
            reverse("legacy_migration:apply", args=[run.pk]),
            {
                "manifest_digest": run.dry_run_manifest_digest,
                "rationale": "Synthetic manager recovery request.",
            },
        )
    assert reapplied.status_code == 200
    assert "withdrawn legacy migrations cannot be applied" in reapplied.content.decode().lower()
    assert PublishedSourceResult.objects.count() == 1


def test_stale_manifest_and_wrong_tenant_reviewer_fail_closed(tmp_path: Path) -> None:
    organization = PartnerOrganization.objects.create(name="Synthetic Stale Show")
    other = PartnerOrganization.objects.create(name="Synthetic Other Reviewer Show")
    manager = _account(
        organization,
        label="legacy-stale-manager",
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
    )
    reviewer = _account(
        organization,
        label="legacy-stale-reviewer",
        role=PrivilegedRoleAssignment.Role.EXPORT_REVIEWER,
    )
    other_reviewer = _account(
        other,
        label="legacy-other-reviewer",
        role=PrivilegedRoleAssignment.Role.EXPORT_REVIEWER,
    )
    manager_client = Client()
    force_login_with_fresh_mfa(manager_client, manager)
    run = _upload(manager_client, organization, _xlsx(), tmp_path)
    with override_settings(MNEMEX_PRIVATE_ARTIFACT_ROOT=tmp_path):
        manager_client.post(
            reverse("legacy_migration:configure", args=[run.pk]),
            {"configuration": _configuration()},
        )
        manager_client.post(reverse("legacy_migration:dry-run", args=[run.pk]))
    run.refresh_from_db()

    reviewer_client = Client()
    force_login_with_fresh_mfa(reviewer_client, reviewer)
    stale = reviewer_client.post(
        reverse("legacy_migration:approve", args=[run.pk]),
        {"manifest_digest": "0" * 64, "rationale": "Synthetic stale digest check."},
    )
    assert stale.status_code == 200
    assert "exact dry-run manifest" in stale.content.decode().lower()
    run.refresh_from_db()
    assert run.state == LegacyMigrationRun.State.DRY_RUN_COMPLETE

    other_client = Client()
    force_login_with_fresh_mfa(other_client, other_reviewer)
    assert (
        other_client.post(
            reverse("legacy_migration:approve", args=[run.pk]),
            {
                "manifest_digest": run.dry_run_manifest_digest,
                "rationale": "Synthetic wrong-tenant attempt.",
            },
        ).status_code
        == 404
    )
