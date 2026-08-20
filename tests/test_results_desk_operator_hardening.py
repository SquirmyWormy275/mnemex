from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path

import pytest
from django.core.exceptions import ValidationError
from django.test import Client
from django.urls import reverse
from openpyxl import Workbook

from mnemex.accounts.models import Account, PrivilegedRoleAssignment
from mnemex.foundation.services import IdempotencyConflict
from mnemex.partners.models import PartnerOrganization
from mnemex.results.models import IngestionRun, MappingTemplate, StagedResult
from mnemex.results.services import ingest_spreadsheet_bytes
from tests.mfa_helpers import enroll_account_mfa, force_login_with_fresh_mfa

pytestmark = pytest.mark.django_db

CANONICAL_HEADERS = [
    "source_result_id",
    "source_revision",
    "source_event_id",
    "event_name",
    "result_date",
    "competitor_name",
    "discipline",
    "score_type",
    "score",
    "heat_id",
    "wood_species",
    "wood_diameter_mm",
    "wood_quality",
]


def _account(email: str) -> Account:
    account = Account.objects.create_user(email=email, password="synthetic-password-123")
    enroll_account_mfa(account)
    PrivilegedRoleAssignment.objects.create(
        account=account,
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
        assigned_by=account,
    )
    return account


def _mapping(organization: PartnerOrganization, actor: Account) -> MappingTemplate:
    return MappingTemplate.objects.create(
        organization=organization,
        name="Canonical",
        version=1,
        field_map={header: header for header in CANONICAL_HEADERS},
        created_by=actor,
    )


def _row(result_id: str, *, score: str = "12.34") -> dict[str, object]:
    return {
        "source_result_id": result_id,
        "source_revision": 1,
        "source_event_id": "synthetic-show-2026",
        "event_name": "Synthetic Summer Show",
        "result_date": "2026-07-18",
        "competitor_name": "Synthetic Competitor",
        "discipline": "UNDERHAND",
        "score_type": "time",
        "score": score,
        "heat_id": "heat-1",
        "wood_species": "white pine",
        "wood_diameter_mm": "325",
        "wood_quality": "8",
    }


def _multisheet_xlsx() -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["This summary is not a results table"])
    results = workbook.create_sheet("Results")
    results.append(CANONICAL_HEADERS)
    results.append([_row("inactive-results")[header] for header in CANONICAL_HEADERS])
    backup = workbook.create_sheet("Backup")
    backup.append(CANONICAL_HEADERS)
    backup.append([_row("backup-results")[header] for header in CANONICAL_HEADERS])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _xlsx_call(
    actor: Account,
    organization: PartnerOrganization,
    mapping: MappingTemplate,
    *,
    operator_key: str,
) -> dict[str, object]:
    return {
        "actor": actor,
        "organization": organization,
        "mapping_template": mapping,
        "operator_key": operator_key,
        "filename": "multi-sheet.xlsx",
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "content": _multisheet_xlsx(),
    }


def test_multisheet_xlsx_requires_an_explicit_worksheet_and_never_uses_active() -> None:
    actor = _account("multi-sheet@mnemex.example.invalid")
    organization = PartnerOrganization.objects.create(name="Synthetic Multi-sheet Show")
    mapping = _mapping(organization, actor)

    with pytest.raises(ValidationError) as error:
        ingest_spreadsheet_bytes(
            **_xlsx_call(actor, organization, mapping, operator_key="multi-no-sheet")
        )

    message = " ".join(error.value.messages)
    assert "multiple worksheets" in message
    assert "Summary" in message
    assert "Results" in message
    assert IngestionRun.objects.count() == 0


def test_explicit_inactive_xlsx_worksheet_publishes_and_persists_coordinates() -> None:
    actor = _account("inactive-sheet@mnemex.example.invalid")
    organization = PartnerOrganization.objects.create(name="Synthetic Inactive Sheet Show")
    mapping = _mapping(organization, actor)

    outcome = ingest_spreadsheet_bytes(
        **_xlsx_call(actor, organization, mapping, operator_key="inactive-sheet"),
        worksheet_name="Results",
    )

    staged = StagedResult.objects.get(run=outcome.run)
    assert outcome.run.published_rows == 1
    assert outcome.run.source_sheet_name == "Results"
    assert staged.source_sheet_name == "Results"
    assert staged.source_row_number == 2
    assert staged.source_result_id == "inactive-results"


def test_invalid_xlsx_worksheet_lists_available_names_and_selection_affects_idempotency() -> None:
    actor = _account("sheet-key@mnemex.example.invalid")
    organization = PartnerOrganization.objects.create(name="Synthetic Sheet Key Show")
    mapping = _mapping(organization, actor)
    call = _xlsx_call(actor, organization, mapping, operator_key="sheet-key")

    with pytest.raises(ValidationError, match="Missing") as error:
        ingest_spreadsheet_bytes(**call, worksheet_name="Missing")
    assert "Results" in " ".join(error.value.messages)

    first = ingest_spreadsheet_bytes(**call, worksheet_name="Results")
    with pytest.raises(IdempotencyConflict):
        ingest_spreadsheet_bytes(**call, worksheet_name="Backup")

    assert first.run.source_sheet_name == "Results"
    assert IngestionRun.objects.count() == 1


def test_csv_persists_physical_source_row_and_no_sheet_name() -> None:
    actor = _account("csv-coordinate@mnemex.example.invalid")
    organization = PartnerOrganization.objects.create(name="Synthetic CSV Coordinate Show")
    mapping = _mapping(organization, actor)
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=CANONICAL_HEADERS)
    writer.writeheader()
    writer.writerow(_row("csv-coordinate"))

    outcome = ingest_spreadsheet_bytes(
        actor=actor,
        organization=organization,
        mapping_template=mapping,
        operator_key="csv-coordinate",
        filename="results.csv",
        content_type="text/csv",
        content=output.getvalue().encode(),
    )

    staged = StagedResult.objects.get(run=outcome.run)
    assert outcome.run.source_sheet_name == ""
    assert staged.source_sheet_name == ""
    assert staged.source_row_number == 2


def test_formula_error_is_quarantined_without_formula_execution() -> None:
    actor = _account("formula-error@mnemex.example.invalid")
    organization = PartnerOrganization.objects.create(name="Synthetic Formula Error Show")
    mapping = _mapping(organization, actor)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(CANONICAL_HEADERS)
    formula_row = _row("formula-error")
    formula_row["score"] = "#DIV/0!"
    sheet.append([formula_row[header] for header in CANONICAL_HEADERS])
    output = BytesIO()
    workbook.save(output)

    outcome = ingest_spreadsheet_bytes(
        actor=actor,
        organization=organization,
        mapping_template=mapping,
        operator_key="formula-error",
        filename="formula.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=output.getvalue(),
    )

    staged = StagedResult.objects.get(run=outcome.run)
    assert staged.outcome == StagedResult.Outcome.QUARANTINED
    assert {error["code"] for error in staged.validation_errors} == {"invalid_number"}


def _manual_post(organization: PartnerOrganization, *, row_count: int) -> dict[str, str]:
    data = {
        "organization": str(organization.pk),
        "operator_key": "more-than-three",
        "rows-TOTAL_FORMS": str(row_count),
        "rows-INITIAL_FORMS": "0",
        "rows-MIN_NUM_FORMS": "0",
        "rows-MAX_NUM_FORMS": "25",
    }
    for index in range(row_count):
        prefix = f"rows-{index}"
        row = _row(f"manual-{index}")
        for field, value in row.items():
            data[f"{prefix}-{field}"] = str(value)
    return data


def test_manual_form_has_local_accessible_add_remove_controls() -> None:
    actor = _account("manual-contract@mnemex.example.invalid")
    organization = PartnerOrganization.objects.create(name="Synthetic Manual Contract Show")
    client = Client()
    force_login_with_fresh_mfa(client, actor)

    response = client.get(reverse("results:manual"))

    assert response.status_code == 200
    html = response.content.decode()
    assert 'id="manual-result-rows"' in html
    assert 'id="add-result-row"' in html
    assert 'type="button"' in html
    assert 'aria-controls="manual-result-rows"' in html
    assert 'id="empty-result-row"' in html
    assert "__prefix__" in html
    assert "results/manual-formset.js" in html
    assert "https://" not in html

    script = (
        Path(__file__).parents[1]
        / "mnemex"
        / "results"
        / "static"
        / "results"
        / "manual-formset.js"
    ).read_text(encoding="utf-8")
    assert "TOTAL_FORMS" in script
    assert "25" in script
    assert "remove" in script.lower()


def test_manual_server_accepts_more_than_three_rows() -> None:
    actor = _account("manual-five@mnemex.example.invalid")
    organization = PartnerOrganization.objects.create(name="Synthetic Manual Five Show")
    client = Client()
    force_login_with_fresh_mfa(client, actor)

    response = client.post(reverse("results:manual"), _manual_post(organization, row_count=5))

    assert response.status_code == 302
    run = IngestionRun.objects.get()
    assert (run.total_rows, run.published_rows) == (5, 5)


def test_issue_csv_is_tenant_scoped_and_contains_stable_recovery_guidance() -> None:
    actor = _account("issue-download@mnemex.example.invalid")
    other = _account("issue-other@mnemex.example.invalid")
    organization = PartnerOrganization.objects.create(name="Synthetic Issue Download Show")
    mapping = _mapping(organization, actor)
    client = Client()
    force_login_with_fresh_mfa(client, actor)
    # Give the second account only an organization-scoped role for another tenant.
    other.privileged_roles.all().delete()
    other_organization = PartnerOrganization.objects.create(name="Other Synthetic Show")
    PrivilegedRoleAssignment.objects.create(
        account=other,
        role=PrivilegedRoleAssignment.Role.RESULTS_MANAGER,
        assigned_by=other,
        scope=PrivilegedRoleAssignment.Scope.ORGANIZATION,
        organization=other_organization,
    )

    valid = _row("issue-valid")
    invalid = _row("issue-invalid", score="not-a-number")
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=CANONICAL_HEADERS)
    writer.writeheader()
    writer.writerows([valid, invalid])
    outcome = ingest_spreadsheet_bytes(
        actor=actor,
        organization=organization,
        mapping_template=mapping,
        operator_key="issue-download",
        filename="issues.csv",
        content_type="text/csv",
        content=output.getvalue().encode(),
    )

    detail = client.get(reverse("results:run-detail", kwargs={"run_id": outcome.run.pk}))
    issue_url = reverse("results:run-issues-csv", kwargs={"run_id": outcome.run.pk})
    assert issue_url in detail.content.decode()

    response = client.get(issue_url)
    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/csv")
    assert response["Content-Disposition"].startswith("attachment;")
    assert response["X-Content-Type-Options"] == "nosniff"
    rows = list(csv.DictReader(StringIO(response.content.decode())))
    assert len(rows) == 1
    assert rows[0]["source_result_id"] == "issue-invalid"
    assert rows[0]["outcome"] == "quarantined"
    assert rows[0]["source_sheet"] == ""
    assert rows[0]["source_row"] == "3"
    assert "score:invalid_number" in rows[0]["validation_reasons"]
    assert "new operator key" in rows[0]["operator_guidance"].lower()
    assert "source revision" in rows[0]["operator_guidance"].lower()

    force_login_with_fresh_mfa(client, other)
    assert client.get(issue_url).status_code == 403


def test_issue_csv_link_is_hidden_when_run_has_no_quarantine_or_conflict() -> None:
    actor = _account("clean-download@mnemex.example.invalid")
    organization = PartnerOrganization.objects.create(name="Synthetic Clean Download Show")
    mapping = _mapping(organization, actor)
    client = Client()
    force_login_with_fresh_mfa(client, actor)
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=CANONICAL_HEADERS)
    writer.writeheader()
    writer.writerow(_row("clean-row"))
    outcome = ingest_spreadsheet_bytes(
        actor=actor,
        organization=organization,
        mapping_template=mapping,
        operator_key="clean-download",
        filename="clean.csv",
        content_type="text/csv",
        content=output.getvalue().encode(),
    )

    detail = client.get(reverse("results:run-detail", kwargs={"run_id": outcome.run.pk}))

    assert "Download rows needing attention" not in detail.content.decode()
